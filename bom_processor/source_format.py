"""BOM source file formats, including macro-enabled workbooks (.xlsm).

A .xlsm is the same Office Open XML spreadsheet package as a .xlsx; only its workbook content type
says "macro-enabled". It is accepted as a BOM source and read exactly like a .xlsx:
- every reader treats the extension .xlsm as .xlsx (reading_extension);
- parsers and the worker receive the same bytes, unchanged and hash-pinned, as input.xlsx
  (staged_input_name), so no parser, reader or worker image needs to know about .xlsm;
- provenance keeps the original file name and records the format (check_source).

Macros are never loaded or run, and external links are never followed or refreshed: openpyxl opens
it read-only with data_only=True (the values Excel stored; no formula is evaluated), keep_vba=False
(the VBA part is never loaded) and keep_links=False. A .xlsm that holds macros is accepted by its
name and read as data only (QQ #99 D1, Nick 2026-10-17); its record says vba_project_present and
carries one information finding, MACROS_IGNORED ("macros ignored"). It is refused when openpyxl
cannot read it or its workbook content type is not a plain or macro-enabled worksheet.
Input: a blob path or extension, and the source bytes. Output: names, extensions and a provenance
dict; failures raise SourceFormatError with a stable code. No I/O beyond reading the given bytes.
Run python -m pytest tests/test_source_format.py.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import PurePosixPath
from typing import Any
from xml.etree import ElementTree

FORMAT_VERSION = "bom-source-format.v1"
MACRO_ENABLED_EXTENSION = ".xlsm"
SOURCE_EXTENSIONS = frozenset({".csv", ".xls", ".xlsx", MACRO_ENABLED_EXTENSION})
# Workbook part content types ([Content_Types].xml) a .xlsm may declare. A .xlsx saved under the
# .xlsm name is still a macro-free worksheet; templates, add-ins and binary workbooks are refused.
WORKBOOK_CONTENT_TYPES = frozenset({
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
})
# The file's own media type, for anything that stores or serves the original .xlsm.
XLSM_MEDIA_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.12"
_CONTENT_TYPES_PART = "[Content_Types].xml"
_TYPES_NAMESPACE = "{http://schemas.openxmlformats.org/package/2006/content-types}"
_DOCTYPE = re.compile(rb"<!DOCTYPE|<!ENTITY", re.I)


MACROS_IGNORED = "MACROS_IGNORED"
MACROS_IGNORED_TEXT = ("macros ignored: the workbook holds macros; they were never loaded or run, "
                       "and only the cell values Excel stored were read")


class SourceFormatError(ValueError):
    """Stable, non-sensitive refusal of a source file's format."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def _extension(value: str) -> str:
    """'.XLSM', 'xlsm', 'doc.xlsm' and 'a/b/doc.xlsm' all give '.xlsm'."""
    text = str(value).casefold()
    if text and "/" not in text and "." not in text.lstrip("."):
        return "." + text.lstrip(".")  # a bare extension, with or without its dot
    return PurePosixPath(text).suffix


def is_macro_enabled(value: str) -> bool:
    """True for a .xlsm path, name or extension (any case)."""
    return _extension(value) == MACRO_ENABLED_EXTENSION


def reading_extension(value: str) -> str:
    """The lower-case dotted extension a source is read as: .xlsm reads as .xlsx; others unchanged."""
    extension = _extension(value)
    return ".xlsx" if extension == MACRO_ENABLED_EXTENSION else extension


def staged_input_name(blob_path: str) -> str:
    """The parser/worker input file name for a source blob: input.xlsx for a .xlsm."""
    return f"input{reading_extension(PurePosixPath(blob_path).suffix)}"


def _workbook_content_type(archive: zipfile.ZipFile) -> str:
    try:
        raw = archive.read(_CONTENT_TYPES_PART)
    except KeyError as exc:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE") from exc
    if _DOCTYPE.search(raw):
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE")
    try:
        root = ElementTree.fromstring(raw)
    except ElementTree.ParseError as exc:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE") from exc
    found = [item.attrib.get("ContentType", "") for item in root.iter(f"{_TYPES_NAMESPACE}Override")
             if item.attrib.get("PartName", "").casefold() == "/xl/workbook.xml"]
    if len(found) != 1:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE")
    if found[0] not in WORKBOOK_CONTENT_TYPES:
        raise SourceFormatError("SOURCE_WORKBOOK_CONTENT_TYPE")
    return found[0]


def check_source(blob_path: str, data: bytes) -> dict[str, Any] | None:
    """Fail closed on a .xlsm that is not a readable worksheet; None for other formats.

    Returns the provenance record for a .xlsm. Every sheet is read through openpyxl, values only,
    so a file a parser could not open is refused here, before any parser or model sees it. Macros
    are not a refusal: they are never loaded, and the record flags them (MACROS_IGNORED).
    """
    if not is_macro_enabled(blob_path):
        return None
    if not isinstance(data, bytes) or not data.startswith(b"PK"):
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE")
    # The survey's OOXML guard: entry, size and ratio bounds, no encryption, no DTDs.
    from .package_survey import _ooxml_guard, _Unreadable
    try:
        has_vba = _ooxml_guard(data)
    except _Unreadable as exc:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = [name.casefold() for name in archive.namelist()]
            # Any VBA part, wherever a writer put it, counts as macros (as the row inventory rules).
            has_vba = has_vba or any("vba" in name for name in names)
            content_type = _workbook_content_type(archive)
    except zipfile.BadZipFile as exc:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE") from exc
    from openpyxl import load_workbook
    book = None
    try:
        # keep_vba=False: no macro part is ever loaded; data_only: stored values, nothing evaluated.
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True,
                             keep_links=False, keep_vba=False)
        if not book.worksheets:
            raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE")
        for sheet in book.worksheets:
            for _ in sheet.iter_rows(values_only=True):
                pass
    except SourceFormatError:
        raise
    except Exception as exc:
        raise SourceFormatError("SOURCE_WORKBOOK_UNREADABLE") from exc
    finally:
        if book is not None:
            book.close()
    findings = ([{"code": MACROS_IGNORED, "severity": "information", "rationale": MACROS_IGNORED_TEXT,
                  "raised_by": "code"}] if has_vba else [])
    return {"version": FORMAT_VERSION, "extension": MACRO_ENABLED_EXTENSION, "read_as": ".xlsx",
            "workbook_content_type": content_type, "vba_project_present": has_vba, "macros_run": False,
            "findings": findings}

"""Hand-written parsers for known BOM layouts, built from one reviewed template (Nick 2026-10-14).

A layout config names, by exact header label, the sheet and the column that holds each canonical
field and each manufacturer/MPN pair. build_parser(config) returns a standalone parse.py: the same
--input/--output/--row-map program an authored candidate is, so it passes the same code policy
(candidate_authoring._validate_code), the same child-process worker and the same schema, semantic
and row-accounting checks, and still needs a person's approval. Nothing here runs a parser.

The template reads the BOM the way the authoring rules describe (candidate_authoring):
- a row with no part number, no quantity and no designator that gives a maker or MPN, directly under
  a part line, is that part's alternate source: its pairs join the line above;
- a row with no designator, no maker/MPN and no quantity (or a printed 0) is a reference line:
  item_category Document, its printed quantity kept (None when blank);
- the one bare board keeps its printed quantity; only when none is printed, and exactly one
  reference line carries the layout's board wording, is it a Part of quantity 1 cited as derived
  with a bom_notes flag;
- a part to place with a blank quantity stays a Part with no quantity, so validation stops; a
  quantity is never counted from designators;
- every cell is written as found; anything a person should check is said in bom_notes.
Columns are found by exact header label; a missing or repeated label fails loudly.

Real layout configs name customer header labels, so they live only in the git-ignored parsers/
directory; tests use synthetic configs. Run: python -m pytest tests/test_fixed_layout_parser.py
"""
from __future__ import annotations

import json
from typing import Any, Mapping

CONFIG_VERSION = "fixed-layout-parser.v1"
FIELD_ROLES = ("item_number", "revision", "item_name", "quantity", "reference_designator",
               "unit_of_measure")
CONTEXT_ROLES = ("category", "comment")
_REQUIRED_ROLES = ("item_number", "item_name", "quantity")


class LayoutConfigError(ValueError):
    """A layout config that cannot describe a parser; the message names the problem."""


def _label(value: Any, where: str) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 200:
        raise LayoutConfigError(f"{where}: header label must be a non-empty string")
    return value


def validate_config(config: Mapping[str, Any]) -> dict[str, Any]:
    """The checked, JSON-ready config; raises LayoutConfigError on anything unclear."""
    if not isinstance(config, Mapping):
        raise LayoutConfigError("config must be an object")
    allowed = {"version", "layout", "sheet", "header_row", "columns", "pairs", "min_pairs", "board_words",
               "ignored_reasons"}
    unknown = set(config) - allowed
    if unknown:
        raise LayoutConfigError(f"unknown config keys: {sorted(unknown)}")
    if config.get("version") != CONFIG_VERSION:
        raise LayoutConfigError(f"version must be {CONFIG_VERSION}")
    layout = config.get("layout")
    if not isinstance(layout, str) or not layout.strip():
        raise LayoutConfigError("layout must name the layout")
    sheet = config.get("sheet")
    if sheet is not None and (not isinstance(sheet, str) or not sheet):
        raise LayoutConfigError("sheet must be a sheet name or null (the one sheet whose headings match)")
    header_row = config.get("header_row")
    if type(header_row) is not int or not 1 <= header_row <= 30:
        raise LayoutConfigError("header_row must be a row number from 1 to 30")
    columns = config.get("columns")
    if not isinstance(columns, Mapping) or set(columns) - set(FIELD_ROLES) - set(CONTEXT_ROLES):
        raise LayoutConfigError(f"columns may name only {FIELD_ROLES + CONTEXT_ROLES}")
    for role in _REQUIRED_ROLES:
        if role not in columns:
            raise LayoutConfigError(f"columns must name {role}")
    checked_columns = {role: _label(label, f"columns.{role}") for role, label in columns.items()}
    pairs = config.get("pairs")
    if not isinstance(pairs, list) or not 1 <= len(pairs) <= 50:
        raise LayoutConfigError("pairs must list 1 to 50 [manufacturer label, MPN label] pairs")
    checked_pairs = []
    for index, pair in enumerate(pairs, 1):
        if not isinstance(pair, list) or len(pair) != 2:
            raise LayoutConfigError(f"pairs[{index}] must be [manufacturer label, MPN label]")
        checked_pairs.append([_label(pair[0], f"pairs[{index}].maker"), _label(pair[1], f"pairs[{index}].mpn")])
    labels = [*checked_columns.values(), *(label for pair in checked_pairs for label in pair)]
    normalized = [" ".join(label.split()).casefold() for label in labels]
    if len(set(normalized)) != len(normalized):
        raise LayoutConfigError("one header label is given two roles")
    # Pairs after the first min_pairs may be missing from a file's headings (a BOM with fewer
    # alternate columns); present ones must come in order with both labels.
    min_pairs = config.get("min_pairs", len(checked_pairs))
    if type(min_pairs) is not int or not 1 <= min_pairs <= len(checked_pairs):
        raise LayoutConfigError("min_pairs must be from 1 to the number of pairs")
    words = config.get("board_words", [])
    if not isinstance(words, list) or any(not isinstance(word, str) or not word.strip() for word in words):
        raise LayoutConfigError("board_words must be a list of words")
    reasons = config.get("ignored_reasons", {})
    if not isinstance(reasons, Mapping) or any(
            not isinstance(key, str) or not isinstance(value, str) or not value.strip() or len(value) > 400
            for key, value in reasons.items()):
        raise LayoutConfigError("ignored_reasons must map header labels to reasons")
    return {"version": CONFIG_VERSION, "layout": layout, "sheet": sheet, "header_row": header_row,
            "columns": checked_columns, "pairs": checked_pairs, "min_pairs": min_pairs,
            "board_words": list(words),
            "ignored_reasons": dict(reasons)}


def build_parser(config: Mapping[str, Any]) -> str:
    """parse.py source for one layout: the template with the checked config embedded as JSON."""
    checked = validate_config(config)
    text = json.dumps(checked, sort_keys=True, ensure_ascii=True)
    return TEMPLATE.replace("CONFIG_JSON = None", "CONFIG_JSON = " + repr(text), 1)


TEMPLATE = r'''"""Hand-written parser for one known BOM layout (bom_processor.fixed_layout_parser template).

Inputs: --input BOM workbook (.xlsx or legacy .xls). Outputs: --output canonical XLSX and, when
asked, --row-map bom-row-map.v2 JSON. The layout config below names every column by header label.
"""
import argparse
import json
from decimal import Decimal, InvalidOperation

import openpyxl
import xlrd
from openpyxl import Workbook

CONFIG_JSON = None
CONFIG = json.loads(CONFIG_JSON)
BASE_COLUMNS = ["line_number", "item_number", "item_name", "revision", "item_category",
                "unit_of_measure", "quantity", "reference_designator", "bom_notes"]
OLE2 = bytes.fromhex("d0cf11e0a1b11ae1")
BOARD_NOTE = ("Selected as this assembly's bare board and given quantity 1 because the BOM prints "
              "no quantity; please confirm")


class LayoutError(ValueError):
    pass


def text(value):
    """A cell as the text a spreadsheet shows; None for blank. Whole floats lose their .0."""
    if value is None or isinstance(value, bool):
        return None if value is None else str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = str(value).strip()
    return result or None


def number(value):
    """(quantity, ok): a finite number, or None when blank; ok is False for unreadable text."""
    shown = text(value)
    if shown is None:
        return None, True
    try:
        parsed = Decimal(shown)
    except (InvalidOperation, ValueError):
        return None, False
    if not parsed.is_finite():
        return None, False
    return (int(parsed) if parsed == parsed.to_integral_value() else float(parsed)), True


def normalize(label):
    return " ".join(str(label).split()).casefold() if label is not None else ""


def read_book(path):
    with open(path, "rb") as handle:
        data = handle.read()
    sheets = []
    if data.startswith(OLE2):
        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            sheets.append((sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)]))
        return sheets
    book = openpyxl.load_workbook(path, data_only=True)
    for sheet in book.worksheets:
        sheets.append((sheet.title, [list(row) for row in sheet.iter_rows(min_row=1, values_only=True)]))
    book.close()
    return sheets


def header_map(row):
    found = {}
    for column, value in enumerate(row, 1):
        key = normalize(value)
        if key:
            found.setdefault(key, []).append(column)
    return found


def present_pairs(found):
    """The configured pairs this heading row shows: the first min_pairs always, later ones in order."""
    pairs = []
    for index, (maker, mpn) in enumerate(CONFIG["pairs"], 1):
        has_maker, has_mpn = normalize(maker) in found, normalize(mpn) in found
        if has_maker and has_mpn:
            if len(pairs) != index - 1:
                return None  # a later pair without an earlier one: not this layout
            pairs.append([maker, mpn])
        elif has_maker or has_mpn or index <= CONFIG["min_pairs"]:
            return None
    return pairs


def locate(sheets):
    """The BOM sheet and every configured label's 1-based column; fails on a missing or repeated label."""
    matches = []
    for name, rows in sheets:
        if CONFIG["sheet"] is not None and name != CONFIG["sheet"]:
            continue
        if len(rows) < CONFIG["header_row"]:
            continue
        found = header_map(rows[CONFIG["header_row"] - 1])
        pairs = present_pairs(found)
        if pairs is not None and all(normalize(label) in found for label in CONFIG["columns"].values()):
            matches.append((name, rows, found, pairs))
    if len(matches) != 1:
        raise LayoutError(f"expected exactly one sheet with this layout's headings, found {len(matches)}")
    name, rows, found, pairs = matches[0]
    wanted = list(CONFIG["columns"].values()) + [label for pair in pairs for label in pair]
    positions = {}
    for label in wanted:
        columns = found[normalize(label)]
        if len(columns) != 1:
            raise LayoutError(f"heading {label!r} appears {len(columns)} times")
        positions[label] = columns[0]
    return name, rows, positions, pairs


def cell(row, column):
    return row[column - 1] if column and column - 1 < len(row) else None


def filled_columns(row):
    return {column for column, value in enumerate(row, 1) if text(value) is not None}


def parse(sheets):
    sheet_name, rows, positions, present = locate(sheets)
    columns = {role: positions[label] for role, label in CONFIG["columns"].items()}
    pair_columns = [(positions[maker], positions[mpn]) for maker, mpn in present]
    board_words = [word.casefold() for word in CONFIG["board_words"]]
    lines = []
    for number_ in range(CONFIG["header_row"] + 1, len(rows) + 1):
        row = rows[number_ - 1]
        if not filled_columns(row):
            continue
        get = lambda role: text(cell(row, columns.get(role)))
        part, ref, name = get("item_number"), get("reference_designator"), get("item_name")
        raw_quantity = cell(row, columns["quantity"])
        quantity, readable = number(raw_quantity)
        pairs = []
        for maker_column, mpn_column in pair_columns:
            maker, mpn = text(cell(row, maker_column)), text(cell(row, mpn_column))
            if maker is not None or mpn is not None:
                pairs.append({"maker": maker, "mpn": mpn, "maker_column": maker_column, "mpn_column": mpn_column})
        if (part is None and text(raw_quantity) is None and ref is None and pairs and lines
                and lines[-1]["category"] == "Part"):
            # An approved alternate source under the part above: its pairs join that line.
            lines[-1]["pairs"].extend(dict(pair, row=number_) for pair in pairs)
            lines[-1]["rows"].append(number_)
            continue
        reference = ref is None and not pairs and readable and (quantity is None or quantity == 0)
        line = {"rows": [number_], "category": "Document" if reference else "Part", "fields": {},
                "values": {}, "pairs": [dict(pair, row=number_) for pair in pairs], "notes": []}
        values, fields = line["values"], line["fields"]
        if part is not None:
            values["item_number"], fields["item_number"] = part, [number_, columns["item_number"]]
        elif not reference and pairs and pairs[0]["mpn"] is not None:
            values["item_number"] = pairs[0]["mpn"]
            fields["item_number"] = [number_, pairs[0]["mpn_column"]]
            line["notes"].append("No part number is printed on this row; the manufacturer part number is "
                                 "used as its item number")
        for role in ("item_name", "revision", "unit_of_measure", "reference_designator"):
            value = get(role)
            if value is not None:
                values[role], fields[role] = value, [number_, columns[role]]
        if quantity is not None:
            values["quantity"], fields["quantity"] = quantity, [number_, columns["quantity"]]
        elif not readable:
            line["notes"].append(f"The quantity cell is not a number ({text(raw_quantity)!r}); left empty "
                                 "for a person")
        comment = get("comment")
        if comment is not None:
            line["notes"].append(f"BOM comment: {comment}")
        clue = " ".join(value for value in (name, get("category"), ref) if value).casefold()
        line["board_clue"] = any(word in clue for word in board_words)
        line["printed_quantity"] = text(raw_quantity) is not None
        lines.append(line)
    choose_board(lines)
    for line in lines:
        describe_pairs(line)
    return sheet_name, rows, lines


def choose_board(lines):
    """The one bare board: kept as printed; derived as 1 only when exactly one candidate prints none."""
    if any(line["board_clue"] and line["category"] == "Part" for line in lines):
        return
    candidates = [line for line in lines if line["board_clue"] and line["category"] == "Document"
                  and not line["printed_quantity"]]
    if len(candidates) != 1:
        return
    board = candidates[0]
    board["category"] = "Part"
    board["values"]["quantity"] = 1
    board["fields"]["quantity"] = {"derived": "the assembly's one bare board; the BOM prints no quantity"}
    board["notes"].append(BOARD_NOTE)


def describe_pairs(line):
    """Say in bom_notes what a person should check about the pairs as written; never rewrite them."""
    pairs = line["pairs"]
    for index, pair in enumerate(pairs, 1):
        if pair["maker"] is None and pair["mpn"] is not None:
            following = pairs[index] if index < len(pairs) else None
            if following is not None and following["row"] == pair["row"] and following["maker"] is not None:
                line["notes"].append(
                    f"Manufacturer pair {index} has a part number ({pair['mpn']!r}) but no manufacturer, and "
                    f"pair {index + 1} on the same row reads {following['maker']!r} / {following['mpn']!r}: "
                    f"the BOM may split one part number across two cells "
                    f"({pair['mpn'] + ',' + (following['mpn'] or '')!r}); please confirm")
            else:
                line["notes"].append(f"Manufacturer pair {index} has a part number but no manufacturer as written")
        elif pair["mpn"] is None:
            line["notes"].append(f"Manufacturer pair {index} has a manufacturer but no part number as written")
        if pair["maker"] and pair["mpn"] and "/" in pair["maker"] and "/" in pair["mpn"]:
            line["notes"].append(f"Manufacturer pair {index} lists several makers and part numbers joined by '/'; "
                                 "kept as written, please split them if they are alternates")


def build(lines):
    width = max([8] + [len(line["pairs"]) for line in lines])
    if width > 50:
        raise LayoutError("a line has more than 50 manufacturer/MPN pairs")
    columns = BASE_COLUMNS + [name for index in range(1, width + 1)
                              for name in (f"manufacturer_{index}", f"manufacturer_item_number_{index}")]
    records = []
    for position, line in enumerate(lines, 1):
        record = dict.fromkeys(columns)
        record.update(line["values"])
        record["line_number"] = position
        record["item_category"] = line["category"]
        for index, pair in enumerate(line["pairs"], 1):
            if pair["maker"] is not None:
                record[f"manufacturer_{index}"] = pair["maker"]
                line["fields"][f"manufacturer_{index}"] = [pair["row"], pair["maker_column"]]
            if pair["mpn"] is not None:
                record[f"manufacturer_item_number_{index}"] = pair["mpn"]
                line["fields"][f"manufacturer_item_number_{index}"] = [pair["row"], pair["mpn_column"]]
        if line["notes"]:
            record["bom_notes"] = "; ".join(line["notes"])
            line["fields"]["bom_notes"] = {"derived": "notes for a person about this line as written"}
        records.append(record)
    return columns, records


def row_map(sheets, sheet_name, rows, lines, positions_by_label):
    mapped = [{"output_row": position + 1, "sheet": sheet_name, "rows": line["rows"], "fields": line["fields"]}
              for position, line in enumerate(lines, 1)]
    skipped = [{"sheet": sheet_name, "rows": [[CONFIG["header_row"], CONFIG["header_row"]]], "reason": "header",
                "rationale": "The column headings of this layout."}]
    if CONFIG["header_row"] > 1:
        skipped.append({"sheet": sheet_name, "rows": [[1, CONFIG["header_row"] - 1]], "reason": "title",
                        "rationale": "Rows above the column headings."})
    for name, other in sheets:
        if name != sheet_name and other:
            skipped.append({"sheet": name, "rows": [[1, len(other)]], "reason": "other_sheet",
                            "rationale": "Not the BOM sheet of this layout; read only the layout's BOM sheet."})
    cited = set()
    for line in lines:
        for link in line["fields"].values():
            if isinstance(link, list):
                cited.add(link[1])
    line_columns = set()
    for line in lines:
        for number_ in line["rows"]:
            line_columns |= filled_columns(rows[number_ - 1])
    labels = {column: label for label, column in positions_by_label.items()}
    headings = rows[CONFIG["header_row"] - 1]
    ignored = []
    for column in sorted(line_columns - cited):
        heading = text(cell(headings, column)) or f"column {column}"
        reason = CONFIG["ignored_reasons"].get(heading) or CONFIG["ignored_reasons"].get(labels.get(column, ""))
        ignored.append({"sheet": sheet_name, "column": column,
                        "reason": reason or f"{heading}: not a canonical field of this layout; read as context only"})
    return {"version": "bom-row-map.v2", "lines": mapped, "skipped": skipped, "ignored_columns": ignored}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--row-map")
    args = parser.parse_args()
    sheets = read_book(args.input)
    sheet_name, rows, lines = parse(sheets)
    if not lines:
        raise LayoutError("the BOM sheet holds no lines")
    _, _, positions, _ = locate(sheets)
    columns, records = build(lines)
    output = Workbook()
    sheet = output.active
    sheet.title = "BOM"
    sheet.append(columns)
    for record in records:
        sheet.append([record[column] for column in columns])
    output.save(args.output)
    check = openpyxl.load_workbook(args.output, read_only=True)
    written = list(check.active.values)
    check.close()
    if list(written[0]) != columns or len(written) != len(records) + 1:
        raise LayoutError("the written workbook does not match the canonical columns and rows")
    if args.row_map:
        with open(args.row_map, "w", encoding="utf-8") as handle:
            json.dump(row_map(sheets, sheet_name, rows, lines, positions), handle)


if __name__ == "__main__":
    main()
'''

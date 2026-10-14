"""Check explicit XLSX document roles without rewriting parsed output.

Inputs: inert workbook bytes and parsed rows. Output: bounded policy evidence,
never source cell contents. Run tests/test_document_roles.py offline.
"""
from __future__ import annotations

import io
import math
import re
import zipfile
from collections import Counter
from typing import Any

from openpyxl import load_workbook

from .canonical_format import pair_indexes

POLICY_VERSION = "xlsx-document-roles.v2"
MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
_DOCUMENT_LABELS = {"document", "documentreference", "drawing", "procedure", "specification"}
_LABEL_HEADERS = {"item", "itemnumber", "customeritem", "category", "itemcategory", "type", "description", "itemname"}
_ID_HEADERS = {"documentid", "documentnumber", "drawingnumber", "procedurenumber", "specificationnumber"}
_OVERLOADED_IDS = {"manufacturerdocumentid", "manufacturerdocumentnumber"}
_REV_HEADERS = {"revision", "rev", "documentrevision", "mpnrevision"}
_NAME_HEADERS = {"description", "itemname", "documentdescription"}
_QTY_HEADERS = {"qty", "quantity"}


def _text(value: Any) -> str:
    if value is None or isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def _quantity(value: Any) -> float | None:
    if not _text(value):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("ambiguous_quantity")
    return float(value)


def check_document_roles(source: bytes, parsed_rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Recognize explicit tabular roles only; unsupported layouts claim no coverage."""
    evidence = {"version": POLICY_VERSION, "valid": False, "result": "source_roles_ambiguous",
                "source_document_rows": 0, "coverage_basis": "xlsx_document_roles"}
    expected: Counter = Counter()
    document_rows = 0
    workbook = None
    try:
        with zipfile.ZipFile(io.BytesIO(source)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("source_limit")
        workbook = load_workbook(io.BytesIO(source), read_only=True, data_only=False, keep_links=False)
        scanned = 0
        for sheet in workbook.worksheets:
            headers: dict[str, list[int]] | None = None
            if sheet.max_column > MAX_COLUMNS or sheet.max_row > MAX_ROWS:
                raise ValueError("source_limit")
            for cells in sheet.iter_rows():
                scanned += 1
                if scanned > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise ValueError("source_limit")
                values = [cell.value for cell in cells]
                keys = [_key(value) for value in values]
                if (set(keys) & (_ID_HEADERS | _OVERLOADED_IDS | {"item", "itemnumber", "customeritem"})
                        and set(keys) & (_NAME_HEADERS | {"category", "itemcategory", "type"})):
                    headers = {}
                    for index, key in enumerate(keys):
                        if key:
                            headers.setdefault(key, []).append(index)
                    continue
                label_indices = ([i for key, indices in headers.items() if key in _LABEL_HEADERS
                                  for i in indices] if headers else range(len(keys)))
                populated = [(i, value) for i, value in enumerate(values) if _text(value)]
                reference_list = (len(populated) == 2 and keys[0] in {"drawing", "procedure", "specification"}
                                  and populated[0][0] == 0 and populated[1][0] == 1)
                if not reference_list and not any(keys[i] in _DOCUMENT_LABELS for i in label_indices):
                    continue
                document_rows += 1
                evidence["source_document_rows"] = document_rows
                # Explicit two-cell reference lists are independent of BOM table headers.
                if reference_list:
                    if any(cells[i].data_type == "f" for i, _ in populated):
                        raise ValueError("source_roles_ambiguous")
                    expected[(_text(values[1]), "", _text(values[0]), None)] += 1
                    continue
                if headers is None:
                    raise ValueError("source_roles_ambiguous")

                def field(roles: set[str], *, required: bool = False) -> tuple[Any, int | None]:
                    indices = [i for key, positions in headers.items() if key in roles for i in positions]
                    if len(indices) > 1 or required and not indices:
                        raise ValueError("source_roles_ambiguous")
                    return (values[indices[0]], indices[0]) if indices else (None, None)

                id_roles = (_ID_HEADERS if set(headers) & _ID_HEADERS else
                            _OVERLOADED_IDS if set(headers) & _OVERLOADED_IDS else {"item", "itemnumber"})
                # Positively labelled documentary rows can overload component
                # columns. Require the exact paired header roles and row label;
                # never infer from an ID prefix or a filename.
                component_reference = ("customeritem" in headers
                    and len(headers["customeritem"]) == 1
                    and keys[headers["customeritem"][0]] == "documentreference"
                    and "manufacturer" in headers and "manufacturerpartnumber" in headers
                    and "description" in headers
                    and keys[headers["description"][0]] in {"drawing", "procedure", "specification", "document"})
                if "customeritem" in headers and len(headers["customeritem"]) != 1:
                    raise ValueError("source_roles_ambiguous")
                if component_reference:
                    id_roles = {"manufacturer"}
                identity, id_index = field(id_roles, required=True)
                revision, rev_index = field({"manufacturerpartnumber"} if component_reference else _REV_HEADERS)
                name, name_index = field(_NAME_HEADERS, required=True)
                quantity, qty_index = field(_QTY_HEADERS)
                for index in (id_index, rev_index, name_index, qty_index):
                    if index is not None and cells[index].data_type == "f":
                        raise ValueError("source_roles_ambiguous")
                if not _text(identity) or _key(identity) in _DOCUMENT_LABELS or not _text(name):
                    raise ValueError("source_roles_ambiguous")
                expected[(_text(identity), _text(revision), _text(name), _quantity(quantity))] += 1
        if not document_rows:
            return {**evidence, "valid": True, "result": "not_applicable",
                    "coverage_basis": "not_applicable"}
        actual: Counter = Counter()
        for row in parsed_rows:
            if _key(row.get("item_category")) != "document":
                continue
            if any(_text(row.get(f"{prefix}_{index}")) for index in pair_indexes(row)
                   for prefix in ("manufacturer", "manufacturer_item_number")):
                return {**evidence, "result": "document_manufacturer_fields"}
            actual[(_text(row.get("item_number")), _text(row.get("revision")),
                    _text(row.get("item_name")), _quantity(row.get("quantity")))] += 1
        return {**evidence, "valid": actual == expected,
                "result": "valid" if actual == expected else "document_fields_or_coverage"}
    except (ValueError, TypeError, zipfile.BadZipFile, KeyError):
        return evidence
    finally:
        if workbook is not None:
            workbook.close()
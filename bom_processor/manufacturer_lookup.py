"""One more parser attempt when makers and MPNs sit on another sheet (QuickQuote A047, Nick 2026-10-14).

A047's BOM lists its lines on one sheet and keeps each part's manufacturer and manufacturer part
number on another sheet keyed by part number. The parser left them out, so validation stopped with
manufacturer_data_missing for a person. When the source itself shows such a sheet (headings that
read as a maker and an MPN column, and a column holding the part numbers of lines without a pair),
the stop allows ONE automatic refinement that brings the values in by lookup link. A second stop
holds for a person as before. Any doubt or error keeps today's hold (fail closed).

Pure detection (lookup_hint, retry_hint, refinement_instructions) plus a one-per-request marker in
the registry (retry_state, claim_retry). Headings are clues, not rules: they decide only whether a
retry is worth one attempt, never what the output holds. Tests: tests/test_manufacturer_lookup.py.
"""
from __future__ import annotations

import io
import json
import re
from typing import Any

HINT_VERSION = "manufacturer-lookup.v1"
MARKER_VERSION = "manufacturer-lookup-retry.v1"
REVIEW_REASON = "manufacturer_data_review_required"
MAX_SHEETS = 5
_MAKER = re.compile(r"\b(?:MANUFACTURERS?|MFRS?|MFG|MAKERS?|VENDORS?)\b")
_PART = re.compile(r"\b(?:PART|PARTS|PN|P N|NUMBER|NUM|NO)\b")
_MPN = re.compile(r"\bMPNS?\b")


def _words(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", " ", str(value).upper()).strip()


def _text(value: Any) -> str:
    # The same comparison row accounting uses for lookup keys.
    return "" if value is None else str(value).strip()


def _unpaired_parts(output: bytes) -> set[str]:
    """Item numbers of Part lines that carry no complete manufacturer/MPN pair."""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(output), read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()
    if not rows:
        return set()
    header = [str(value) if value is not None else "" for value in rows[0]]
    index = {name: position for position, name in enumerate(header)}
    pairs = [(index[name], index.get(name.replace("manufacturer_", "manufacturer_item_number_", 1)))
             for name in header if re.fullmatch(r"manufacturer_\d+", name)]
    parts = set()
    for row in rows[1:]:
        get = lambda position: row[position] if position is not None and position < len(row) else None  # noqa: E731
        if _words(get(index.get("item_category"))) != "PART":
            continue
        if any(_text(get(maker)) and _text(get(mpn)) for maker, mpn in pairs):
            continue
        item = _text(get(index.get("item_number")))
        if item:
            parts.add(item)
    return parts


def lookup_hint(source: bytes, extension: str, output: bytes, line_sheets: set[str]) -> dict[str, Any] | None:
    """Sheets other than the lines' own that hold makers and MPNs for unpaired parts, or None."""
    from .row_accounting import source_inventory
    unpaired = _unpaired_parts(output)
    if not unpaired:
        return None
    inventory = source_inventory(source, extension)
    found = []
    for sheet in inventory.sheets:
        if sheet in line_sheets:
            continue
        rows: dict[int, dict[int, Any]] = {}
        for (name, row, column), value in inventory.cells.items():
            if name == sheet and _text(value):
                rows.setdefault(row, {})[column] = value
        header_row = next((row for row in sorted(rows)
                           if sum(isinstance(value, str) for value in rows[row].values()) >= 2), None)
        if header_row is None:
            continue
        makers, mpns = [], []
        for column, value in sorted(rows[header_row].items()):
            words = _words(value)
            if _MPN.search(words) or (_MAKER.search(words) and _PART.search(words)):
                mpns.append(column)
            elif _MAKER.search(words):
                makers.append(column)
        if not makers or not mpns:
            continue
        best = None
        for key in sorted({column for row in rows.values() for column in row} - set(makers) - set(mpns)):
            hits = {_text(cells[key]) for row, cells in rows.items()
                    if row > header_row and key in cells and _text(cells[key]) in unpaired
                    and any(_text(cells.get(column)) for column in mpns)}
            if hits and (best is None or len(hits) > best[1]):
                best = (key, len(hits))
        if best is not None:
            found.append({"sheet": sheet, "header_row": header_row, "key_column": best[0],
                          "manufacturer_columns": makers, "mpn_columns": mpns, "parts": best[1]})
    if not found:
        return None
    return {"version": HINT_VERSION, "sheets": found[:MAX_SHEETS]}


def _columns(value: Any) -> bool:
    return isinstance(value, list) and 0 < len(value) <= 20 and all(type(item) is int and item >= 1 for item in value)


def retry_hint(event: dict[str, Any]) -> dict[str, Any] | None:
    """The validated hint of a manufacturer stop that may be retried, else None."""
    hint = event.get("manufacturer_lookup")
    if event.get("review_reason") != REVIEW_REASON or not isinstance(hint, dict):
        return None
    sheets = hint.get("sheets")
    if (set(hint) != {"version", "sheets"} or hint["version"] != HINT_VERSION or not isinstance(sheets, list)
            or not 0 < len(sheets) <= MAX_SHEETS):
        return None
    for item in sheets:
        if (not isinstance(item, dict)
                or set(item) != {"sheet", "header_row", "key_column", "manufacturer_columns", "mpn_columns", "parts"}
                or not isinstance(item["sheet"], str) or not 0 < len(item["sheet"]) <= 128
                or not all(type(item[key]) is int and item[key] >= 1 for key in ("header_row", "key_column", "parts"))
                or not _columns(item["manufacturer_columns"]) or not _columns(item["mpn_columns"])):
            return None
    return hint


def refinement_instructions(hint: dict[str, Any]) -> str:
    return ("MANUFACTURER LOOKUP REFINEMENT. The previous parser (supplied as previous_parser_code) wrote part "
            "lines without any complete manufacturer/MPN pair, but the source keeps manufacturer names and "
            "manufacturer part numbers on another sheet, one row per part number. Sheets, with the 1-based "
            "heading row, the column holding part numbers (key_column) and the maker and MPN columns: "
            + json.dumps(hint["sheets"], ensure_ascii=False, separators=(",", ":"))
            + ". Correct previous_parser_code so every line whose part number equals a key_column cell on such a "
            "sheet carries each of those rows' manufacturer and MPN as one pair (a second row for the same part is "
            "an alternate pair of the same line), citing both halves as {\"lookup\": [sheet, row, column], "
            "\"key\": key_column}; rows a line now cites by lookup are no longer listed in skipped. A part with no "
            "row there keeps its manufacturer fields empty; never invent a pair. Keep everything that already "
            "passed unchanged and return the complete replacement.")


def _marker_path(request_id: str) -> str:
    return f"candidate-automation/{request_id}/manufacturer-lookup-retry.json"


def retry_state(storage: Any, request_id: str, candidate_id: str) -> str:
    """eligible (no retry yet), dispatched (the retry is this candidate's) or spent; spent on any doubt."""
    from .core import BlobNotFoundError
    try:
        marker = json.loads(storage.download_registry(_marker_path(request_id), 10_000).decode("utf-8"))
    except (BlobNotFoundError, FileNotFoundError):
        return "eligible"
    except Exception:
        return "spent"
    if (isinstance(marker, dict) and marker.get("version") == MARKER_VERSION
            and marker.get("request_id") == request_id and marker.get("source_candidate") == candidate_id):
        return "dispatched"
    return "spent"


def claim_retry(storage: Any, request_id: str, candidate_id: str, validation_id: str) -> bool:
    """Write the one-per-request marker; True when this candidate holds it."""
    from .core import BlobConflictError, _json_bytes
    body = {"version": MARKER_VERSION, "request_id": request_id, "source_candidate": candidate_id,
            "validation_id": validation_id}
    try:
        storage.upload_registry_immutable(_marker_path(request_id), _json_bytes(body), "application/json")
    except BlobConflictError:
        pass
    return retry_state(storage, request_id, candidate_id) == "dispatched"

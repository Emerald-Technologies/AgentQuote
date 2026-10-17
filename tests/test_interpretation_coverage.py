"""Synthetic offline coverage fixtures; never injected into live agent requests."""
import copy
import io
import json
from hashlib import sha256

from openpyxl import Workbook
import pytest

from bom_processor.core import CANONICAL_COLUMNS
from bom_processor.interpretation_inventory import EXTRACTION_VERSION, InventoryError, extract_inventory
from bom_processor.interpretation_coverage import CoverageError, verify_coverage


def xlsx(rows, title="BOM", extra=None):
    book = Workbook(); sheet = book.active; sheet.title = title
    for row in rows:
        sheet.append(row)
    if extra is not None:
        book.create_sheet("Notes").append(extra)
    result = io.BytesIO(); book.save(result); book.close()
    return result.getvalue()


def ref(row, column):
    return {"sheet": "BOM", "row": row, "column": column}


def fixture(document_quantity=None):
    # The user's layout: two physical rows of alternatives per component.
    source = xlsx([
        ["Part", "Qty", "MPN", "MFG"],
        ["A", 1, "abc123", "jad"],
        [None, None, "xyz123", "oks"],
        ["D", 5, "iwejf123", "wwo9ekf"],
        [None, None, "oiekjfr", "okewf"],
        ["DRAW-1", document_quantity, "Specification"],
    ])
    records = []
    rows = [list(CANONICAL_COLUMNS)]
    for output_row, source_row, item in [(2, 2, "A"), (3, 4, "D")]:
        values = dict.fromkeys(CANONICAL_COLUMNS)
        values.update(line_number=output_row-1, item_number=item, item_name=item,
                      item_category="Part", quantity=1 if item == "A" else 5,
                      manufacturer_1="jad" if item == "A" else "wwo9ekf",
                      manufacturer_item_number_1="abc123" if item == "A" else "iwejf123",
                      manufacturer_2="oks" if item == "A" else "okewf",
                      manufacturer_item_number_2="xyz123" if item == "A" else "oiekjfr")
        rows.append([values[key] for key in CANONICAL_COLUMNS])
        fields = {"item_number": ref(source_row, 1), "item_name": ref(source_row, 1),
                  "quantity": ref(source_row, 2), "manufacturer_1": ref(source_row, 4),
                  "manufacturer_item_number_1": ref(source_row, 3),
                  "manufacturer_2": ref(source_row+1, 4), "manufacturer_item_number_2": ref(source_row+1, 3)}
        records.append({"id": item, "output_row": output_row, "category": "Part", "fields": fields})
    doc = dict.fromkeys(CANONICAL_COLUMNS)
    doc.update(line_number=3, item_number="DRAW-1", item_name="Specification", item_category="Document",
               quantity=document_quantity)
    rows.append([doc[key] for key in CANONICAL_COLUMNS])
    fields = {"item_number": ref(6, 1), "item_name": ref(6, 3)}
    if document_quantity is not None:
        fields["quantity"] = ref(6, 2)
    records.append({"id": "doc", "output_row": 4, "category": "Document", "fields": fields})
    output = xlsx(rows, title="Parsed")
    inventory = extract_inventory(source, ".xlsx")
    coverage = []
    for (_, row, col) in inventory.cells:
        ids = [] if row == 1 else ["A"] if row in (2, 3) else ["D"] if row in (4, 5) else ["doc"]
        role = "structural" if row == 1 else "alternate" if row in (3, 5) else "document" if row == 6 else "component"
        coverage.append({"cell": ref(row, col), "disposition": role, "record_ids": ids,
                         "rationale": "Synthetic agent interpretation for offline regression."})
    proposal = {"version": "bom-interpretation.v1",
                "source": {"sha256": sha256(source).hexdigest(), "size_bytes": len(source),
                           "extraction_version": EXTRACTION_VERSION},
                "output": {"sha256": sha256(output).hexdigest(), "size_bytes": len(output), "sheet": "Parsed"},
                "records": records, "coverage": coverage, "findings": []}
    return source, output, proposal


def verify(source, output, proposal):
    return verify_coverage(source, ".xlsx", output, json.dumps(proposal).encode())


@pytest.mark.parametrize("quantity", [None, 0, 3])
def test_vertical_alternatives_and_optional_document_quantity(quantity):
    source, output, proposal = fixture(quantity)
    result = verify(source, output, proposal)
    assert result["output_records"] == 3
    assert result["provenance_verified"] is True
    assert result["has_blocking_findings"] is False
    assert result["automatic_acceptance_authorized"] is False
    assert result["semantic_acceptance"] == "not_evaluated"


@pytest.mark.parametrize("mutation, code", [
    (lambda p: p.update(extra=True), "SCHEMA_KEYS_INVALID"),
    (lambda p: p["source"].update(sha256="0"*64), "SOURCE_BINDING_INVALID"),
    (lambda p: p["output"].update(sha256="0"*64), "OUTPUT_BINDING_INVALID"),
    (lambda p: p["coverage"].pop(), "SOURCE_COVERAGE_INCOMPLETE"),
    (lambda p: p["coverage"].append(copy.deepcopy(p["coverage"][0])), "SOURCE_COVERAGE_DUPLICATE_OR_UNKNOWN"),
    (lambda p: p["records"][0]["fields"]["quantity"].update(row=4), "SOURCE_VALUE_CHANGED"),
    (lambda p: p["records"][0]["fields"].pop("manufacturer_2"), "OUTPUT_FIELD_COVERAGE_INVALID"),
    (lambda p: p["records"][0]["fields"]["manufacturer_item_number_2"].update(row=2), "SOURCE_VALUE_CHANGED"),
    (lambda p: p["records"][0]["fields"]["item_number"].update(row=90), "SOURCE_CELL_MISSING"),
])
def test_tampering_rejected(mutation, code):
    source, output, proposal = fixture(); mutation(proposal)
    with pytest.raises(CoverageError, match=code):
        verify(source, output, proposal)


def test_duplicate_json_keys_and_oversized_sidecar_rejected():
    source, output, proposal = fixture()
    data = json.dumps(proposal).replace('"version":', '"version":"duplicate","version":', 1).encode()
    with pytest.raises(CoverageError, match="DUPLICATE_JSON_KEY"):
        verify_coverage(source, ".xlsx", output, data)
    with pytest.raises(CoverageError, match="SIDECAR_SIZE_LIMIT"):
        verify_coverage(source, ".xlsx", output, b" " * (1024*1024+1))


def test_unresolved_source_unit_and_findings_do_not_authorize_acceptance():
    source, output, proposal = fixture()
    proposal["coverage"][0].update(disposition="unresolved", rationale="Cannot establish meaning.")
    proposal["findings"] = [{"code": "RELATION_UNCERTAIN", "severity": "blocking",
                             "rationale": "Interpretation needs review.", "cells": [ref(1, 1)]}]
    result = verify(source, output, proposal)
    assert result["has_blocking_findings"] is True
    assert result["automatic_acceptance_authorized"] is False


def test_inventory_includes_every_sheet_and_never_evaluates_formula():
    source = xlsx([["Part", "Qty"], ["A", 1]], extra=["Supporting evidence"])
    inventory = extract_inventory(source, ".xlsx")
    assert inventory.sheets == ("BOM", "Notes")
    assert inventory.cells[("Notes", 1, 1)] == "Supporting evidence"
    # openpyxl writes no stored result, so the value is unknowable without evaluation.
    with pytest.raises(InventoryError, match="SOURCE_FORMULA_RESULT_MISSING"):
        extract_inventory(xlsx([["=1+1"]]), ".xlsx")
    # A macro-free .xlsm is read as the .xlsx it is (2026-10-14); other formats are still refused.
    assert extract_inventory(source, ".xlsm").cells == inventory.cells
    with pytest.raises(InventoryError, match="SOURCE_FORMAT_UNSUPPORTED"):
        extract_inventory(source, ".xlsb")


def test_html_xls_is_inert_and_merged_anchors_not_filled_down():
    source = b"<html><body><table><tr><td>A</td><td>1</td></tr><tr><td></td><td>alternate</td></tr></table></body></html>"
    inventory = extract_inventory(source, ".xls")
    assert ("HTML_TABLE_1", 2, 1) not in inventory.cells
    assert inventory.cells[("HTML_TABLE_1", 2, 2)] == "alternate"
    book = Workbook(); sheet = book.active; sheet.merge_cells("A1:A2"); sheet["A1"] = "A"
    stream = io.BytesIO(); book.save(stream); book.close()
    inventory = extract_inventory(stream.getvalue(), ".xlsx")
    assert inventory.merged_ranges == (("Sheet", "A1:A2"),)
    assert ("Sheet", 2, 1) not in inventory.cells


@pytest.mark.parametrize("mutation, code", [
    (lambda p: p["source"].update(size_bytes=float(p["source"]["size_bytes"])), "BINDING_TYPE_INVALID"),
    (lambda p: p["records"][0].update(category=[]), "RECORD_INVALID"),
    (lambda p: p["findings"].append({"code": "TEST", "severity": [], "rationale": "Reason", "cells": []}), "FINDING_INVALID"),
])
def test_strict_binding_and_enum_types(mutation, code):
    source, output, proposal = fixture(); mutation(proposal)
    with pytest.raises(CoverageError, match=code):
        verify(source, output, proposal)


def test_coordinated_dropped_alternative_is_a_blocking_gap():
    from openpyxl import load_workbook
    source, output, proposal = fixture()
    book = load_workbook(io.BytesIO(output)); sheet = book["Parsed"]
    sheet.cell(2, 12).value = None; sheet.cell(2, 13).value = None
    stream = io.BytesIO(); book.save(stream); book.close(); output = stream.getvalue()
    proposal["records"][0]["fields"].pop("manufacturer_2")
    proposal["records"][0]["fields"].pop("manufacturer_item_number_2")
    proposal["output"].update(sha256=sha256(output).hexdigest(), size_bytes=len(output))
    result = verify(source, output, proposal)
    assert result["unrepresented_semantic_cells"] == 2
    assert result["has_blocking_findings"] is True


def test_sparse_dimensions_rejected_before_workbook_materialization():
    book = Workbook(); sheet = book.active
    sheet["A1"] = "small"; sheet["SR100000"] = "large rectangle"
    stream = io.BytesIO(); book.save(stream); book.close()
    with pytest.raises(InventoryError, match="SOURCE_TRAVERSAL_LIMIT"):
        extract_inventory(stream.getvalue(), ".xlsx")


def _with_stored_results(source, results):
    """Add Excel-style stored results (<v>) to formula cells, as Excel does on save."""
    import re
    import zipfile
    archive = zipfile.ZipFile(io.BytesIO(source))
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as output:
        for name in archive.namelist():
            data = archive.read(name)
            if name.startswith("xl/worksheets/"):
                text = data.decode()
                for ref, (kind, value) in results.items():
                    text = re.sub(rf'<c r="{ref}"([^>]*)><f>([^<]*)</f><v\s*/>',
                                  rf'<c r="{ref}"\1 t="{kind}"><f>\2</f><v>{value}</v>', text)
                    text = re.sub(rf'<c r="{ref}"([^>]*)><f>([^<]*)</f></c>',
                                  rf'<c r="{ref}"\1 t="{kind}"><f>\2</f><v>{value}</v></c>', text)
                data = text.encode()
            output.writestr(name, data)
    return stream.getvalue()


def test_formula_cells_use_stored_result_and_are_recorded():
    source = _with_stored_results(xlsx([["Part", "Qty"], ["A", "=1+1"]]), {"B2": ("n", "2")})
    inventory = extract_inventory(source, ".xlsx")
    assert inventory.cells[("BOM", 2, 2)] == 2
    assert inventory.formula_cells == frozenset({("BOM", 2, 2)})
    assert inventory.formula_cells_known is True


def test_formula_with_stored_error_is_refused():
    source = _with_stored_results(xlsx([["Part", "Qty"], ["A", "=1/0"]]), {"B2": ("e", "#DIV/0!")})
    with pytest.raises(InventoryError, match="SOURCE_ERROR_CELL"):
        extract_inventory(source, ".xlsx")


def test_parser_output_with_formula_is_refused():
    source, output, proposal = fixture()
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(output)); book["Parsed"].cell(2, 7).value = "=0+1"
    stream = io.BytesIO(); book.save(stream); book.close()
    output = _with_stored_results(stream.getvalue(), {"G2": ("n", "1")})
    proposal["output"].update(sha256=sha256(output).hexdigest(), size_bytes=len(output))
    with pytest.raises(CoverageError, match="OUTPUT_FORMULA_PRESENT"):
        verify(source, output, proposal)


def test_invalid_binary_xls_is_refused():
    with pytest.raises(InventoryError, match="SOURCE_WORKBOOK_INVALID"):
        extract_inventory(b"\xd0\xcf\x11\xe0" + b"unused", ".xls")
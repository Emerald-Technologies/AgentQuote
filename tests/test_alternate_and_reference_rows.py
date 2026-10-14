"""Rows without a quantity of their own (QQ A052, Nick 2026-10-14).

A052's parsers spent all six attempts on two row kinds that have no quantity:
- kind 1, an approved alternate source: no item or part number and no quantity, only a maker and an
  MPN, directly under a real part. It joins that part's line as another pair, never a new part;
- kind 2, reference rows at the top of the BOM naming the assembly and its documents, with no
  designator and no quantity. They are listed as Document lines with quantity None (not priced). The
  one bare board among them is bought: with no printed quantity it becomes a Part of quantity 1 with a
  bom_notes flag for a person, and validation accepts that derived quantity for one line per assembly
  group only (each refusal is tested below);
- kind 3 must not change: a real part with a designator and a blank quantity still stops validation,
  and its quantity is never derived from the designators.
The authoring rules describe these as clues; here a real parser (child process) that follows them runs
against synthetic workbooks through the offline validation path, and a schema failure now tells the
next attempt which rule it broke. Synthetic data only; no Azure or model calls.
Run: python -m pytest tests/test_alternate_and_reference_rows.py
"""
import hashlib
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from bom_processor import InMemoryStorage
from bom_processor.candidate_authoring import (ALTERNATE_ROW_RULE, BARE_BOARD_RULE, BLANK_PART_QUANTITY_RULE,
                                               REFERENCE_ROW_RULE, CandidateAuthoringService)
from bom_processor.candidate_validation import CandidateValidationError, CandidateValidationService
from bom_processor.core import CANONICAL_COLUMNS, _json_bytes
from legacy_xls import legacy_xls
from test_combined_bom_legacy_xls import LocalWorker

# A parser written the way the rules ask, for a layout given by CONFIG (1-based sheet columns).
# alternates="pair" puts a maker/MPN-only row on the line above; "line" writes it as its own Part
# (the A052 mistake). Reference rows become Documents with quantity None, except the bare board
# (quantity 1, derived and flagged); a blank part quantity stays blank.
PARSER = '''import argparse
import json
import openpyxl
import xlrd
from openpyxl import Workbook
CONFIG = None
COLUMNS = None
NOTE = "Selected as this assembly's bare board and given quantity 1 because the BOM prints no quantity; please confirm"
p = argparse.ArgumentParser()
p.add_argument("--input", required=True)
p.add_argument("--output", required=True)
p.add_argument("--row-map")
args = p.parse_args()
with open(args.input, "rb") as handle:
    data = handle.read()
if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
    sheet = xlrd.open_workbook(file_contents=data).sheet_by_name(CONFIG["sheet"])
    rows = [sheet.row_values(index) for index in range(sheet.nrows)]
else:
    book = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    rows = [list(row) for row in book[CONFIG["sheet"]].iter_rows(values_only=True)]
    book.close()
cols = CONFIG["cols"]
def get(row, role):
    column = cols.get(role)
    value = row[column - 1] if column and column - 1 < len(row) else None
    if isinstance(value, str):
        value = value.strip()
    return None if value in (None, "") else value
def text(value):
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
records, lines = [], []
for index, row in enumerate(rows, 1):
    if index <= CONFIG["header_row"] or not any(v not in (None, "") for v in row):
        continue
    pn, qty, ref = get(row, "pn"), get(row, "qty"), get(row, "ref")
    maker, mpn, desc = get(row, "maker"), get(row, "mpn"), get(row, "desc")
    if pn is None and qty is None and (maker or mpn) and records and CONFIG["alternates"] == "pair":
        record, line = records[-1], lines[-1]
        slot = 1 + max([0] + [n for n in range(1, 9) if record["manufacturer_" + str(n)]])
        record["manufacturer_" + str(slot)] = text(maker)
        record["manufacturer_item_number_" + str(slot)] = text(mpn)
        line["fields"]["manufacturer_" + str(slot)] = [index, cols["maker"]]
        line["fields"]["manufacturer_item_number_" + str(slot)] = [index, cols["mpn"]]
        line["rows"].append(index)
        continue
    record = dict.fromkeys(COLUMNS)
    fields = {}
    reference = pn is not None and qty is None and ref is None and maker is None and mpn is None
    # The test double picks the board by its wording; the real author weighs several clues.
    board = reference and "FABRICATION" in str(desc)
    record.update(line_number=len(records) + 1, item_category="Document" if reference and not board else "Part",
                  item_number=text(pn if pn is not None else mpn), item_name=text(desc), quantity=1 if board else qty)
    fields["item_number"] = [index, cols["pn"] if pn is not None else cols["mpn"]]
    fields["item_name"] = [index, cols["desc"]]
    if qty is not None:
        fields["quantity"] = [index, cols["qty"]]
    if ref is not None:
        record["reference_designator"] = text(ref)
        fields["reference_designator"] = [index, cols["ref"]]
    if not reference and (maker or mpn):
        record["manufacturer_1"] = text(maker) if maker else None
        record["manufacturer_item_number_1"] = text(mpn) if mpn else None
        if maker:
            fields["manufacturer_1"] = [index, cols["maker"]]
        if mpn:
            fields["manufacturer_item_number_1"] = [index, cols["mpn"]]
    if board:
        record["bom_notes"] = NOTE
        fields["bom_notes"] = {"derived": "flag for a person"}
        fields["quantity"] = {"derived": "the assembly's one bare board; the BOM prints no quantity"}
    records.append(record)
    lines.append({"output_row": len(records) + 1, "sheet": CONFIG["sheet"], "rows": [index], "fields": fields})
output = Workbook()
out = output.active
out.append(COLUMNS)
for record in records:
    out.append([record[column] for column in COLUMNS])
output.save(args.output)
if args.row_map:
    skipped = [{"sheet": CONFIG["sheet"], "rows": [[CONFIG["header_row"], CONFIG["header_row"]]],
                "reason": "header", "rationale": "Column headings."}]
    ignored = [{"sheet": CONFIG["sheet"], "column": column, "reason": "Item position only."}
               for column in CONFIG["ignored"]]
    with open(args.row_map, "w", encoding="utf-8") as handle:
        json.dump({"version": "bom-row-map.v2", "lines": lines, "skipped": skipped, "ignored_columns": ignored}, handle)
'''

# Kind 1, the A052 MOD BOM shape (legacy .xls): item 1's alternate maker sits on the next row with no
# item, part number or quantity.
KIND1 = [["Item", "IPN", "Quantity", "Manufacturer", "MPN", "REFDES-List", "Desc-1"],
         [1, "P-0001", 4, "MAKER A", "MPN-A-10", "C1,C2,C3,C4", "CAP CER 10PF 50V"],
         [None, None, None, "MAKER B", "MPN-B-10", None, "CAP CER 10PF 50V"],
         [2, "P-0002", 2, "MAKER C", "MPN-C-1K", "R1,R2", "RES 1K 0603"]]
KIND1_CONFIG = {"sheet": "MOD BOM", "header_row": 1, "ignored": [1],
                "cols": {"pn": 2, "qty": 3, "maker": 4, "mpn": 5, "ref": 6, "desc": 7}}
# Kind 2, the A052 "additional documents" shape (.xlsx): a blank row, the headings, a blank row, the
# assembly and its documents with only a part number and description, a blank row, then the parts.
KIND2 = [[None] * 7,
         ["ITEM", "QTY", "REF", "PART NO", "MFGR", "MFGR PN", "DESCRIPTION"],
         [None] * 7,
         [None, None, None, "ASM-0200", None, None, "PCB ASSY, TEST BOARD"],
         [None, None, None, "ART-0200", None, None, "PCB ARTWORK, TEST BOARD"],
         [None, None, None, "SCH-0200", None, None, "PCB SCHEMATIC, TEST BOARD"],
         [None, None, None, "FAB-0200", None, None, "PCB FABRICATION, TEST BOARD"],
         [None] * 7,
         [1, 3, "C1,C2,C3", "P-0010", "MAKER A", "MPN-A-1U", "CAP CER 1UF"],
         [2, 1, "U1", "P-0011", "MAKER D", "MPN-D-REG", "IC REG LDO"]]
KIND2_CONFIG = {"sheet": "ASSY_REV02", "header_row": 2, "ignored": [1],
                "cols": {"qty": 2, "ref": 3, "pn": 4, "maker": 5, "mpn": 6, "desc": 7}}
# Kind 3: a real part with a designator and a maker/MPN but a blank quantity (an A052 shape).
KIND3 = [["Level", "Ref Des", "Part Number", "Description", "Quantity", "MFG PN1", "MFG NAME1"],
         ["1", "C1, C8", "P-0020", "CAP CER 1UF", 2, "MPN-A-1U", "MAKER A"],
         ["1", "C70", "P-0021", "CAP CER 100PF", None, "MPN-K-100P", "MAKER K"]]
KIND3_CONFIG = {"sheet": "PARTS", "header_row": 1, "ignored": [1],
                "cols": {"ref": 2, "pn": 3, "desc": 4, "qty": 5, "mpn": 6, "maker": 7}}


class Model:
    """Returns the configured parser and keeps the rules it was given."""
    def __init__(self, parser):
        self.parser, self.rules = parser, ()

    def generate(self, prompt, source_bytes, prior_code=None, failed_rule_id=None):
        self.rules = prompt.rules
        return self.parser


def xlsx(sheet, rows):
    book = Workbook()
    book.active.title = sheet
    for row in rows:
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def parser(config, alternates="pair"):
    return (PARSER.replace("CONFIG = None", "CONFIG = " + repr({**config, "alternates": alternates}))
            .replace("COLUMNS = None", "COLUMNS = " + repr(list(CANONICAL_COLUMNS))))


def run(source, extension, code, request="alt-ref-rows"):
    store = InMemoryStorage()
    path = f"customer/rfq-12345678-1234-4234-8234-123456789012/package-1/original-boms/bom{extension}"
    store.sources[path] = source
    store.registry[f"candidates/requests/{request}.json"] = _json_bytes({
        "version": "candidate-request.v2", "request_id": request, "status": "requested", "parser_key": "Parser1",
        "customer_key": "customer", "profile_hash": "a" * 64, "profile": {"extension": extension},
        "source_evidence": [{"document_id": "doc-1", "blob_path": path, "size_bytes": len(source),
                             "sha256": hashlib.sha256(source).hexdigest()}]})
    model = Model(code)
    authored = CandidateAuthoringService(request, store, model).author()
    service = CandidateValidationService(authored.request_id, authored.candidate_id, store, LocalWorker(store))
    return store, model, authored, service


def output_rows(store, authored):
    [key] = [key for key in store.candidates
             if key.startswith(f"candidates/{authored.request_id}/{authored.candidate_id}/") and key.endswith("output.xlsx")]
    book = load_workbook(io.BytesIO(store.candidates[key]), read_only=True)
    values = list(book.active.values)
    book.close()
    return [dict(zip(values[0], row)) for row in values[1:]]


def failure_event(store, authored):
    [key] = [key for key in store.candidates
             if key.startswith(f"candidates/events/{authored.request_id}/{authored.candidate_id}/validation-failed-")]
    return json.loads(store.candidates[key])


def test_the_rules_describe_alternate_reference_and_blank_quantity_rows_as_clues():
    _, model, _, _ = run(legacy_xls({"MOD BOM": KIND1}), ".xls", parser(KIND1_CONFIG))
    for rule in (ALTERNATE_ROW_RULE, REFERENCE_ROW_RULE, BLANK_PART_QUANTITY_RULE):
        assert rule in model.rules
    assert "clues a person reads together, not a rule" in ALTERNATE_ROW_RULE
    assert "add its maker and MPN as the next manufacturer/MPN pair of the line above" in ALTERNATE_ROW_RULE
    assert "item_category='Document' line" in REFERENCE_ROW_RULE and "quantity None" in REFERENCE_ROW_RULE
    assert "never drop such a row and never give it a quantity" in REFERENCE_ROW_RULE
    assert "given quantity 1 because the BOM prints no quantity" in BARE_BOARD_RULE
    assert "Never mark more than one line per assembly as its board" in BARE_BOARD_RULE
    assert "leave it for a person" in BARE_BOARD_RULE and "clues read together" in BARE_BOARD_RULE
    assert BARE_BOARD_RULE in model.rules
    assert "never count designators" in BLANK_PART_QUANTITY_RULE
    assert "never call it a Document" in BLANK_PART_QUANTITY_RULE


def test_kind1_an_alternate_source_row_joins_the_part_above_as_a_second_pair():
    store, _, authored, service = run(legacy_xls({"MOD BOM": KIND1}), ".xls", parser(KIND1_CONFIG))
    assert service.validate().status == "awaiting_approval"
    rows = output_rows(store, authored)
    assert [(row["item_number"], row["quantity"]) for row in rows] == [("P-0001", 4), ("P-0002", 2)]
    first = rows[0]
    assert (first["manufacturer_1"], first["manufacturer_item_number_1"]) == ("MAKER A", "MPN-A-10")
    assert (first["manufacturer_2"], first["manufacturer_item_number_2"]) == ("MAKER B", "MPN-B-10")


def test_kind1_an_alternate_written_as_its_own_part_fails_and_the_feedback_names_the_rule():
    store, _, authored, service = run(legacy_xls({"MOD BOM": KIND1}), ".xls", parser(KIND1_CONFIG, "line"))
    with pytest.raises(CandidateValidationError) as failure:
        service.validate()
    assert failure.value.code == "SCHEMA_INVALID"
    event = failure_event(store, authored)
    assert event["repairable"] is True
    message = event["failure_message"]
    assert message.startswith("canonical_schema_invalid: quantity invalid: ")
    assert "that part's alternate: its pair goes on the line above" in message
    assert "never derive it from designators" in message


def test_kind2_reference_rows_are_documents_and_the_one_board_gets_quantity_1_with_a_flag():
    store, _, authored, service = run(xlsx("ASSY_REV02", KIND2), ".xlsx", parser(KIND2_CONFIG))
    assert service.validate().status == "awaiting_approval"
    rows = output_rows(store, authored)
    assert [(row["item_number"], row["item_category"], row["quantity"]) for row in rows] == [
        ("ASM-0200", "Document", None), ("ART-0200", "Document", None), ("SCH-0200", "Document", None),
        ("FAB-0200", "Part", 1), ("P-0010", "Part", 3), ("P-0011", "Part", 1)]
    # The bare board is bought: one line, quantity 1, flagged for a person; the documents stay blank.
    assert rows[3]["bom_notes"].startswith("Selected as this assembly's bare board and given quantity 1")
    assert all(row["bom_notes"] is None for row in rows[:3])
    assert all(row["manufacturer_1"] is None for row in rows[:4])


def test_kind3_a_real_part_with_a_blank_quantity_still_stops_and_gets_no_quantity():
    store, _, authored, service = run(xlsx("PARTS", KIND3), ".xlsx", parser(KIND3_CONFIG))
    with pytest.raises(CandidateValidationError) as failure:
        service.validate()
    assert failure.value.code == "SCHEMA_INVALID"
    rows = output_rows(store, authored)
    assert [(row["item_number"], row["item_category"], row["quantity"]) for row in rows] == [
        ("P-0020", "Part", 2), ("P-0021", "Part", None)]
    assert "quantity invalid" in failure_event(store, authored)["failure_message"]


def test_an_unknown_schema_failure_keeps_the_bare_code():
    from bom_processor.candidate_validation import _schema_feedback
    assert _schema_feedback(None) == "canonical_schema_invalid"
    assert _schema_feedback("something else") == "canonical_schema_invalid"


# Row accounting for the derived board quantity (Nick 2026-10-14): synthetic sources, direct checks.
NOTE = "Selected as the bare board; quantity 1 because the BOM prints none"
DERIVED = {"quantity": {"derived": "one bare board"}, "bom_notes": {"derived": "flag"}}
PART_LINE = ({"item_number": "P-1", "item_name": "CAP 1UF", "item_category": "Part", "quantity": 2,
              "reference_designator": "C1,C2", "manufacturer_1": "MAKER A", "manufacturer_item_number_1": "MPN-A"},
             5, {"item_number": [5, 1], "item_name": [5, 2], "quantity": [5, 3], "reference_designator": [5, 4],
                 "manufacturer_1": [5, 5], "manufacturer_item_number_1": [5, 6]})
BOARD_SOURCE = [["Part", "Desc", "Qty", "Ref", "MFG", "MPN"],
                ["ASM-1", "PCB ASSY, TEST", None, None, None, None],
                ["ART-1", "PCB ARTWORK, TEST", None, None, None, None],
                ["FAB-1", "PCB FABRICATION, TEST", None, None, None, None],
                ["P-1", "CAP 1UF", 2, "C1,C2", "MAKER A", "MPN-A"]]


def document(part, name, row):
    return ({"item_number": part, "item_name": name, "item_category": "Document"}, row,
            {"item_number": [row, 1], "item_name": [row, 2]})


def board(part, name, row, values=None, fields=None):
    line = {"item_number": part, "item_name": name, "item_category": "Part", "quantity": 1, "bom_notes": NOTE}
    return ({**line, **(values or {})}, row, {"item_number": [row, 1], "item_name": [row, 2], **(fields or DERIVED)})


def accounting(lines, source=None, parents=None):
    """Row accounting of one synthetic BOM: lines are (output values, source row, fields)."""
    from bom_processor.canonical_format import canonical_columns
    from bom_processor.row_accounting import verify_row_map
    columns = canonical_columns(8)
    book = [columns] + [[{**dict.fromkeys(columns), "line_number": n, **values}[c] for c in columns]
                        for n, (values, _, _) in enumerate(lines, 1)]
    mapped = []
    for n, (_, row, fields) in enumerate(lines, 1):
        line = {"output_row": n + 1, "sheet": "BOM", "rows": [row], "fields": fields}
        if parents is not None:
            line["parent"] = {"line": parents[n - 1], "from": [row, 1], "by": "level"}
        mapped.append(line)
    row_map = {"version": "bom-row-map.v2", "ignored_columns": [], "lines": mapped,
               "skipped": [{"sheet": "BOM", "rows": [[1, 1]], "reason": "header", "rationale": "Headings."}]}
    return verify_row_map(xlsx("BOM", source or BOARD_SOURCE), ".xlsx", xlsx("Parsed", book),
                          json.dumps(row_map).encode())


def test_the_one_board_may_take_a_derived_quantity_of_1_with_a_flag():
    result = accounting([document("ASM-1", "PCB ASSY, TEST", 2), document("ART-1", "PCB ARTWORK, TEST", 3),
                         board("FAB-1", "PCB FABRICATION, TEST", 4), PART_LINE])
    assert result["valid"] is True, result["failure_codes"]


@pytest.mark.parametrize("values,fields,cells", [
    ({"quantity": 2}, None, {}),                                   # a derived quantity other than 1
    ({"bom_notes": None}, {"quantity": {"derived": "one bare board"}}, {}),   # no flag for a person
    ({"item_category": "Document"}, None, {}),                     # only a Part line is the board
    # A placed part with a blank quantity (kind 3) is never given one: a designator, or a maker/MPN.
    ({"reference_designator": "C9"}, {**DERIVED, "reference_designator": [4, 4]}, {3: "C9"}),
    ({"manufacturer_1": "MAKER Z", "manufacturer_item_number_1": "MPN-Z"},
     {**DERIVED, "manufacturer_1": [4, 5], "manufacturer_item_number_1": [4, 6]}, {4: "MAKER Z", 5: "MPN-Z"}),
])
def test_a_derived_board_quantity_is_refused_outside_its_narrow_shape(values, fields, cells):
    source = [list(row) for row in BOARD_SOURCE]
    for column, value in cells.items():
        source[3][column] = value
    codes = accounting([document("ASM-1", "PCB ASSY, TEST", 2), document("ART-1", "PCB ARTWORK, TEST", 3),
                        board("FAB-1", "PCB FABRICATION, TEST", 4, values, fields), PART_LINE],
                       source=source)["failure_codes"]
    assert codes == {"BOARD_QUANTITY_DERIVATION_INVALID": 1}, codes


def test_a_derived_quantity_where_the_bom_printed_a_value_is_refused():
    printed = [list(row) for row in BOARD_SOURCE]
    printed[3][2] = 0
    codes = accounting([document("ASM-1", "PCB ASSY, TEST", 2), document("ART-1", "PCB ARTWORK, TEST", 3),
                        board("FAB-1", "PCB FABRICATION, TEST", 4), PART_LINE], source=printed)["failure_codes"]
    assert "BOARD_QUANTITY_PRINTED" in codes, codes


def test_a_second_derived_board_in_the_same_assembly_is_refused():
    codes = accounting([document("ASM-1", "PCB ASSY, TEST", 2), board("ART-1", "PCB ARTWORK, TEST", 3),
                        board("FAB-1", "PCB FABRICATION, TEST", 4), PART_LINE])["failure_codes"]
    assert "BOARD_QUANTITY_REPEATED" in codes, codes


def test_two_assemblies_each_take_their_own_one_board():
    """A two-level BOM: each sub-assembly's group has its own board line with no printed quantity."""
    source = [["Level", "Part", "Desc", "Qty", "Ref", "MFG", "MPN"],
              [1, "SUB-A", "SUB ASSY A", 1, None, None, None],
              [2, "FAB-A", "PCB FABRICATION A", None, None, None, None],
              [2, "PA-1", "CAP A", 2, "C1,C2", "MAKER A", "MPN-A"],
              [1, "SUB-B", "SUB ASSY B", 1, None, None, None],
              [2, "FAB-B", "PCB FABRICATION B", None, None, None, None],
              [2, "PB-1", "RES B", 1, "R1", "MAKER B", "MPN-B"]]

    def line(row, part, name, values, fields):
        return ({"item_number": part, "item_name": name, "item_category": "Part", **values}, row,
                {"item_number": [row, 2], "item_name": [row, 3], **fields})

    def part(row, number, name, quantity, ref, maker, mpn):
        return line(row, number, name, {"quantity": quantity, "reference_designator": ref, "manufacturer_1": maker,
                                         "manufacturer_item_number_1": mpn},
                    {"quantity": [row, 4], "reference_designator": [row, 5], "manufacturer_1": [row, 6],
                     "manufacturer_item_number_1": [row, 7]})

    lines = [line(2, "SUB-A", "SUB ASSY A", {"quantity": 1}, {"quantity": [2, 4]}),
             line(3, "FAB-A", "PCB FABRICATION A", {"quantity": 1, "bom_notes": NOTE}, DERIVED),
             part(4, "PA-1", "CAP A", 2, "C1,C2", "MAKER A", "MPN-A"),
             line(5, "SUB-B", "SUB ASSY B", {"quantity": 1}, {"quantity": [5, 4]}),
             line(6, "FAB-B", "PCB FABRICATION B", {"quantity": 1, "bom_notes": NOTE}, DERIVED),
             part(7, "PB-1", "RES B", 1, "R1", "MAKER B", "MPN-B")]
    # Output rows 2..7 match source rows 2..7; each board's parent is its own sub-assembly's line.
    result = accounting(lines, source=source, parents=[None, 2, 2, None, 5, 5])
    assert result["valid"] is True, (result["failure_codes"], result["failed_rows"])
    # Both boards claimed for the first sub-assembly: two boards in one group are refused.
    claimed = accounting(lines, source=source, parents=[None, 2, 2, None, 2, 5])
    assert "BOARD_QUANTITY_REPEATED" in claimed["failure_codes"], claimed["failure_codes"]

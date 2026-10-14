"""Offline run of multi-level BOMs through the whole package pipeline (sub-assembly plan unit 7).

QuickQuote Docs/SUBASSEMBLY_BOM_PLAN_2026-10-12.md, Nick 2026-10-12. Three synthetic workbooks hold
the same tree in the three in-file forms: a Level column, dotted numbering and a parent part-number
column. Makers carry a customer code prefix ("63 ANALOG DEVICES INC.") and an alternate sits on the
following row. A fake model returns the parser; the parser runs as a real child process and writes
row map v2 with parent claims; validation, review, approval and resolution run unchanged. The
published structure files must match one answer key. No Azure/model/CQ calls; synthetic data only.
Run: pytest tests/test_subassembly_package_integration.py
"""
import hashlib
import io
import json
from datetime import datetime, timezone
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from bom_processor import InMemoryStorage, ParentPackageService, ProcessingRequest
from bom_processor.assembly_association import VERSION as ASSOCIATION_VERSION
from bom_processor.candidate_authoring import CandidateAuthoringService
from bom_processor.candidate_validation import CandidateValidationService
from bom_processor.core import CANONICAL_COLUMNS, OUTPUT_CONTENT_TYPE, _json_bytes
from bom_processor.document_role import VERSION as ROLE_VERSION
from bom_processor.foundry_author import FoundryModelAuthor
from bom_processor.job_orchestrator import StartResult
from worker.runtime import Limits, run_job

RFQ = "12345678-1234-4234-8234-123456789012"
REQUEST = "multi-level-demo"
MAKER = "63 ANALOG DEVICES INC."
# (structure value per form: level, dotted number, parent part), part, description, quantity, maker, MPN, alternate MPN.
TREE = [
    ((1, "1", None), "PCA-1", "PCB ASSY MAIN", 2, None, None, None),
    ((2, "1.1", "PCA-1"), "PCB-1", "BARE PCB", 1, None, None, None),
    ((2, "1.2", "PCA-1"), "U1", "IC ADC", 3, MAKER, "AD7606BSTZ", "AD7606BSTZ-RL"),
    ((1, "2", None), "CBL-1", "CABLE ASSY FAN", 1, None, None, None),
    ((2, "2.1", "CBL-1"), "CON-1", "CONNECTOR", 2, None, None, None),
    ((1, "3", None), "SCR-1", "SCREW", 4, None, None, None),
]
# Answer key by part: (quantity per parent, quantity per quoted assembly, procurement).
ANSWERS = {"PCA-1": ("2", "2", "build"), "PCB-1": ("1", "2", "buy"), "U1": ("3", "6", "buy"),
           "CBL-1": ("1", "1", "buy"), "CON-1": ("2", "2", "reference"), "SCR-1": ("4", "4", "buy")}
# form: (assembly, structure heading, whether the quoted assembly is its own row)
FORMS = {"level": ("ASM-L", "Level", True), "number": ("ASM-D", "Find", False), "parent": ("ASM-P", "Parent", True)}

PARSER = '''import argparse
import json
import re
from openpyxl import load_workbook, Workbook
p = argparse.ArgumentParser()
p.add_argument("--input", required=True)
p.add_argument("--output", required=True)
p.add_argument("--row-map", required=True)
args = p.parse_args()
source = load_workbook(args.input, data_only=False)
columns = COLUMNS
KINDS = {"Level": "level", "Find": "number", "Parent": "parent"}
records = []
skipped = []
def skip(sheet, number, reason, rationale):
    skipped.append({"sheet": sheet.title, "rows": [[number, number]], "reason": reason, "rationale": rationale})
def maker_value(text):
    # A customer code in front of the maker name is not part of the name; cite it as a code link.
    match = re.fullmatch(r"\\d{1,3}\\s+(\\S.*)", text)
    return (match.group(1), True) if match else (text, False)
for sheet in source:
    headers = None
    current = None
    for number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True), 1):
        values = list(cells)
        if not any(value is not None for value in values):
            continue
        if "Part" in values and "Qty" in values:
            headers = {str(value): i for i, value in enumerate(values) if value is not None}
            skip(sheet, number, "header", "Column headings.")
            continue
        if headers is None:
            skip(sheet, number, "title", "Assembly declarations above the header.")
            continue
        def get(key):
            index = headers.get(key)
            return values[index] if index is not None and index < len(values) else None
        def at(key, code=False):
            return [number, headers[key] + 1] + (["code"] if code else [])
        item, quantity, maker, mpn = get("Part"), get("Qty"), get("Maker"), get("MPN")
        if item is None and quantity is None and maker is not None and mpn is not None:
            if current is None:
                raise ValueError("alternate without parent")
            slot = next(i for i in range(1, 9) if current.get("manufacturer_" + str(i)) is None)
            name, coded = maker_value(str(maker))
            current["manufacturer_" + str(slot)] = name
            current["manufacturer_item_number_" + str(slot)] = str(mpn)
            current["_map"]["rows"].append(number)
            current["_map"]["fields"]["manufacturer_" + str(slot)] = at("Maker", coded)
            current["_map"]["fields"]["manufacturer_item_number_" + str(slot)] = at("MPN")
        elif item is not None and quantity is not None:
            heading = next((key for key in KINDS if key in headers), None)
            record = dict.fromkeys(columns)
            record.update(item_number=str(item), item_name=str(get("Desc")), item_category="Part", quantity=quantity)
            fields = {"item_number": at("Part"), "item_name": at("Desc"), "quantity": at("Qty")}
            if get("Rev") is not None:
                record["revision"] = str(get("Rev"))
                fields["revision"] = at("Rev")
            if maker is not None and mpn is not None:
                name, coded = maker_value(str(maker))
                record.update(manufacturer_1=name, manufacturer_item_number_1=str(mpn))
                fields.update(manufacturer_1=at("Maker", coded), manufacturer_item_number_1=at("MPN"))
            record["_map"] = {"sheet": sheet.title, "rows": [number], "fields": fields}
            record["_structure"] = (KINDS[heading], get(heading), at(heading)) if heading else None
            records.append(record)
            current = record
        else:
            raise ValueError("unresolved source row")
source.close()
# Parent of each line (output row), from the structure column; quantities stay as printed.
for index, record in enumerate(records):
    if record["_structure"] is None:
        continue  # a flat BOM states no parents
    kind, value, cell = record["_structure"]
    earlier = list(reversed(list(enumerate(records[:index]))))
    if kind == "level":
        parent = next((i for i, r in earlier if r["_structure"][1] < value), None)
    elif kind == "number":
        head = str(value).rsplit(".", 1)[0] if "." in str(value) else None
        parent = next((i for i, r in earlier if str(r["_structure"][1]) == head), None) if head else None
    else:
        parent = next((i for i, r in earlier if value and r["item_number"].casefold() == str(value).casefold()), None)
    record["_map"]["parent"] = {"line": None if parent is None else parent + 2, "from": cell, "by": kind}
output = Workbook()
sheet = output.active
sheet.append(columns)
lines = []
for i, record in enumerate(records, 1):
    record["line_number"] = i
    record.pop("_structure")
    mapped = record.pop("_map")
    lines.append({"output_row": i + 1, **mapped})
    sheet.append([record[column] for column in columns])
output.save(args.output)
with open(args.row_map, "w", encoding="utf-8") as handle:
    json.dump({"version": "bom-row-map.v2", "lines": lines, "skipped": skipped, "ignored_columns": []}, handle)
output.close()
'''.replace("columns = COLUMNS", "columns = " + repr(list(CANONICAL_COLUMNS)))


def workbook(form):
    assembly, heading, own_row = FORMS[form]
    book = Workbook()
    sheet = book.active
    sheet.title = "Customer"
    sheet.append(["Assembly", assembly])
    sheet.append(["Revision", "A"])
    sheet.append(["Part", "Desc", "Qty", "Maker", "MPN", heading])
    if own_row:
        sheet.append([assembly, "TOP ASSEMBLY", 1, None, None, 0 if form == "level" else None])
    for values, part, desc, qty, maker, mpn, alternate in TREE:
        structure = {"level": values[0], "number": values[1], "parent": values[2] or assembly}[form]
        sheet.append([part, desc, qty, maker, mpn, structure])
        if alternate:
            sheet.append([None, None, None, maker, alternate, None])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def book(assembly, revision, header, rows):
    """A workbook in the fixture layout: assembly and revision above a header row at row 3."""
    data = Workbook()
    sheet = data.active
    sheet.title = "Customer"
    sheet.append(["Assembly", assembly])
    sheet.append(["Revision", revision])
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    data.save(stream)
    data.close()
    return stream.getvalue()


def execute(parser=PARSER, through="resolve", books=None):
    """books: [(assembly, revision, workbook bytes)]; the three forms by default."""
    store = InMemoryStorage()
    sources, associations, roles = [], {}, {}
    books = books or [(FORMS[form][0], "A", workbook(form)) for form in FORMS]
    for index, (assembly, revision, source) in enumerate(books, 1):
        digest = hashlib.sha256(source).hexdigest()
        path = f"customer/rfq-{RFQ}/package-1/original-boms/attachment-{index}.xlsx"
        store.sources[path] = source
        sources.append({"document_id": f"doc-{index}", "blob_path": path, "sha256": digest, "size_bytes": len(source)})
        associations[digest] = {"version": ASSOCIATION_VERSION, "source_sha256": digest,
            "name": {"value": assembly, "cell": {"sheet": "Customer", "row": 1, "column": 2}},
            "revision": {"value": revision, "cell": {"sheet": "Customer", "row": 2, "column": 2}},
            "rationale": "Synthetic: assembly and revision declared above the BOM.", "findings": []}
        roles[digest] = {"version": ROLE_VERSION, "source_sha256": digest, "role": "bom",
            "rationale": "Synthetic: one row per part with quantity and a structure column.",
            "evidence": [{"value": "Part", "cell": {"sheet": "Customer", "row": 3, "column": 1}}], "findings": []}
    store.registry[f"candidates/requests/{REQUEST}.json"] = _json_bytes({
        "version": "candidate-request.v2", "request_id": REQUEST, "status": "requested", "parser_key": "Parser1",
        "customer_key": "customer", "profile_hash": "a" * 64, "profile": {"extension": ".xlsx"},
        "source_evidence": sources})

    class Agent:
        def __init__(self, **kwargs):
            assert kwargs["tools"] == []
        async def run(self, text, *, options):
            return SimpleNamespace(output_text=json.dumps({
                "parser_code": parser, "document_roles": roles, "assembly_associations": associations}))

    authored = CandidateAuthoringService(REQUEST, store, FoundryModelAuthor(agent_factory=Agent),
                                         source_associations=True).author()

    class LocalWorker:
        def start(self, item):
            ran = run_job(job=item.job_id, parser_bytes=store.candidates[item.parser_blob],
                input_bytes=store.sources[item.source_blob], input_name=item.input_name,
                output_name=item.output_name, output_policy="opaque", return_output=True,
                limits=Limits(timeout_seconds=15), row_map_name=item.row_map_name)
            result, code, output = ran[:3]
            assert code == 0, (result["failure"], result.get("stderr_tail"))
            # Storage transport only; output and row map come from the real child process.
            store.candidates[item.output_blob] = output
            store.candidates[item.row_map_blob] = ran[3]
            store.candidates[item.result_blob] = _json_bytes({**result, "contract_version": "remote-worker.v3",
                "input": {**result["input"], "verified": True}, "parser": {**result["parser"], "verified": True},
                "output": {**result["output"], "content_type": OUTPUT_CONTENT_TYPE},
                "row_map": {**result["row_map"], "content_type": "application/json"}})
            return StartResult(item.job_id, "offline-child-process", datetime.now(timezone.utc))

    validation = CandidateValidationService(authored.request_id, authored.candidate_id, store, LocalWorker()).validate()
    if through == "validation":
        return store, validation
    assert validation.status == "awaiting_approval", validation
    packages = ParentPackageService(store)
    plan = packages.plan(ProcessingRequest.from_dict({"job_id": "multi-level-demo", "customer_key": "customer",
        "parser_key": "Parser1", "rfq_uuid": RFQ, "package_version": "1", "documents": sources}),
        approved_assignments=[], pending_assignments=[{
            "request_id": authored.request_id, "document_ids": [item["document_id"] for item in sources]}])
    review = packages.review(plan["plan_id"])
    assert review["approval_ready"], review
    approval = packages.approve(plan_id=plan["plan_id"], review_manifest_sha256=review["manifest_sha256"],
                                approver="offline-test")
    resolved = packages.resolve(plan_id=plan["plan_id"], approval_id=approval["approval_id"])
    return store, packages, plan["plan_id"], approval["approval_id"], resolved


def published_rows(store, result):
    raw = store.download_output(result["output"]["blob_path"], 1_000_000)
    book = load_workbook(io.BytesIO(raw), read_only=True)
    values = list(book.active.values)
    book.close()
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_all_three_forms_resolve_to_the_same_rolled_structure():
    store, packages, plan_id, approval_id, resolved = execute()
    assert resolved["version"] == "resolved-output-manifest.v2"
    by_name = {result["assembly_identity"]["name"]: result for result in resolved["results"]}
    assert set(by_name) == {"asm-l", "asm-d", "asm-p"}
    for form, (assembly, _, own_row) in FORMS.items():
        result = by_name[assembly.casefold()]
        structure = json.loads(packages.resolved_structure(plan_id, approval_id, result["group_id"]))
        assert structure["proven"] is True, (form, structure["findings"])
        lines = {line["line_key"].split("/")[-1]: line for line in structure["lines"]}
        got = {part: (line["quantity_per_parent"], line["quantity"], line["procurement"])
               for part, line in lines.items() if part != assembly}
        assert got == ANSWERS, form
        if own_row:
            assert lines[assembly]["procurement_basis"] == {"kind": "quoted_assembly"}
            assert lines["PCA-1"]["line_key"] == f"{assembly}/PCA-1"
        assert lines["U1"]["line_key"].endswith("PCA-1/U1")
        assert {line["line_key"] for line in structure["lines"] if line["is_subassembly"]} == {
            lines["PCA-1"]["line_key"], lines["CBL-1"]["line_key"]}
        # Every structure line names the published row it describes.
        rows = {row["line_number"]: row for row in published_rows(store, result)}
        for part, line in lines.items():
            assert rows[line["line_number"]]["item_number"].casefold() == part.casefold()
        # Maker code dropped, alternate kept on the same line, quantity as printed (per parent).
        u1 = rows[lines["U1"]["line_number"]]
        assert (u1["manufacturer_1"].casefold(), u1["manufacturer_2"].casefold()) == ("analog devices inc.",) * 2
        assert u1["manufacturer_item_number_2"].casefold() == "ad7606bstz-rl"
        assert u1["quantity"] == 3
    # Replays are stable.
    assert packages.resolve(plan_id=plan_id, approval_id=approval_id) == resolved


def test_a_parser_that_misstates_parents_is_stopped_at_validation():
    # Siblings at the same level claimed as parent and child: the row check works out the real
    # parents from the Level cells and fails the claim, so nothing reaches review.
    wrong = PARSER.replace('r["_structure"][1] < value', 'r["_structure"][1] <= value')
    assert wrong != PARSER
    store, validation = execute(wrong, through="validation")
    assert validation.status != "awaiting_approval"
    evidence = json.loads(next(value for key, value in store.candidates.items()
                               if key.endswith(f"/validation/{validation.validation_id}/evidence.json")))
    text = json.dumps(evidence)
    assert "STRUCTURE_PARENT_MISMATCH" in text

"""A046's shape through the whole package pipeline: a legacy .xls holding both quoted boards (Nick 2026-10-14).

QuickQuote A046 (a copy of A045) found no BOM for either board. Its one BOM is a legacy binary .xls
listing both boards under a Parent Part column, each board with its own artwork and fabrication
drawing rows (Documents), and the manufacturer part numbers on a separate sheet keyed by part number.
There is no assembly association. Validation used to refuse the output because the Document rows
name several items (ASSEMBLY_IDENTITY_CONFLICT), so the split never ran; the author was also never
told to read a legacy .xls with xlrd. Here the authoring rules carry the .xls rule, a real parser
(child process) reads the file with xlrd and states parents, validation passes without a single
identity, and review splits the file into one group per board. Synthetic data only; no Azure/model
calls. Run: pytest tests/test_combined_bom_legacy_xls.py
"""
import hashlib
import io
import json
from datetime import datetime, timezone

from openpyxl import load_workbook

from bom_processor import InMemoryStorage, ParentPackageService, ProcessingRequest
from bom_processor import bom_package_choice
from bom_processor.candidate_authoring import CandidateAuthoringService
from bom_processor.candidate_validation import CandidateValidationService
from bom_processor.core import CANONICAL_COLUMNS, OUTPUT_CONTENT_TYPE, _json_bytes
from bom_processor.job_orchestrator import StartResult
from legacy_xls import legacy_xls
from worker.runtime import Limits, run_job

RFQ = "12345678-1234-4234-8234-123456789012"
REQUEST = "combined-legacy-xls"
A1 = "aaaaaaaa-1111-4111-8111-111111111111"
A2 = "bbbbbbbb-2222-4222-8222-222222222222"
REQUESTED = [{"assembly_id": A1, "part_number": "0100-04931", "revision": "02"},
             {"assembly_id": A2, "part_number": "0100-04932", "revision": "02"}]
# The A045 layout: ORIGINAL, Parent Part, Part Number, Description, Quantity, Level; a board's own row
# names itself as its parent.
BOM = [["ORIGINAL", "Parent Part", "Part Number", "Description", "Quantity", "Level"],
       ["10055369-002", "0100-04931", "0100-04931", "PCB ASSY, BCB", 1, "1"],
       ["10055369-002", "0100-04931", "0120-05191", "PCB ARTWORK", 1, "2"],
       ["10055369-002", "0100-04931", "0110-04455", "PCB FABRICATION", 1, "2"],
       ["10055369-002", "0100-04931", "0630-03236", "CAP CER 33PF", 2, "2"],
       ["10055369-002", "0100-04932", "0100-04932", "PCB ASSY, BCB", 1, "1"],
       ["10055369-002", "0100-04932", "0120-05192", "PCB ARTWORK", 1, "2"],
       ["10055369-002", "0100-04932", "0110-04456", "PCB FABRICATION", 1, "2"],
       ["10055369-002", "0100-04932", "0630-03323", "2225 SIZE CAP", 3, "2"]]
MFG = [["Original", "Part Number", "Description", "Manufacturer Part Number", "Manufacturer Name"],
       ["10055369-002", "0630-03236", "CAP CER 33PF", "C0805C330J5GACTU", "KEMET"],
       ["10055369-002", "0630-03323", "2225 SIZE CAP", "C2225C102KGRACTU", "KEMET"]]

PARSER = '''import argparse
import json
import xlrd
from openpyxl import Workbook
p = argparse.ArgumentParser()
p.add_argument("--input", required=True)
p.add_argument("--output", required=True)
p.add_argument("--row-map", required=True)
args = p.parse_args()
with open(args.input, "rb") as handle:
    book = xlrd.open_workbook(file_contents=handle.read())
columns = COLUMNS
bom, mfg = book.sheet_by_name("BOM"), book.sheet_by_name("MFG")
makers = {}
for index in range(1, mfg.nrows):
    part = str(mfg.cell_value(index, 1)).strip()
    makers.setdefault(part, []).append((index + 1, str(mfg.cell_value(index, 4)), str(mfg.cell_value(index, 3))))
records, lines = [], []
for index in range(1, bom.nrows):
    number = index + 1
    parent_part, part, name = (str(bom.cell_value(index, column)).strip() for column in (1, 2, 3))
    record = dict.fromkeys(columns)
    drawing = "ARTWORK" in name or "FABRICATION" in name
    record.update(line_number=len(records) + 1, item_number=part, item_name=name,
                  item_category="Document" if drawing else "Part", quantity=bom.cell_value(index, 4))
    fields = {"item_number": [number, 3], "item_name": [number, 4], "quantity": [number, 5]}
    for slot, (row, maker, mpn) in enumerate(makers.get(part, []), 1):
        record["manufacturer_" + str(slot)] = maker
        record["manufacturer_item_number_" + str(slot)] = mpn
        fields["manufacturer_" + str(slot)] = {"lookup": ["MFG", row, 5], "key": 2}
        fields["manufacturer_item_number_" + str(slot)] = {"lookup": ["MFG", row, 4], "key": 2}
    earlier = [position for position, item in enumerate(records) if item["item_number"] == parent_part]
    records.append(record)
    lines.append({"output_row": len(records) + 1, "sheet": "BOM", "rows": [number], "fields": fields,
                  "parent": {"line": earlier[-1] + 2 if earlier else None, "from": [number, 2], "by": "parent"}})
output = Workbook()
sheet = output.active
sheet.append(columns)
for record in records:
    sheet.append([record[column] for column in columns])
output.save(args.output)
skipped = [{"sheet": name, "rows": [[1, 1]], "reason": "header", "rationale": "Column headings."} for name in ("BOM", "MFG")]
ignored = [{"sheet": "BOM", "column": column, "reason": "Assembly copy and level."} for column in (1, 6)]
with open(args.row_map, "w", encoding="utf-8") as handle:
    json.dump({"version": "bom-row-map.v2", "lines": lines, "skipped": skipped, "ignored_columns": ignored}, handle)
'''.replace("columns = COLUMNS", "columns = " + repr(list(CANONICAL_COLUMNS)))


class Model:
    """Returns the parser and keeps the rules it was given."""
    def __init__(self):
        self.rules = ()

    def generate(self, prompt, source_bytes, prior_code=None, failed_rule_id=None):
        self.rules = prompt.rules
        return PARSER


class LocalWorker:
    def __init__(self, store):
        self.store = store

    def start(self, item):
        store = self.store
        ran = run_job(job=item.job_id, parser_bytes=store.candidates[item.parser_blob],
                      input_bytes=store.sources[item.source_blob], input_name=item.input_name,
                      output_name=item.output_name, output_policy="opaque", return_output=True,
                      limits=Limits(timeout_seconds=15), row_map_name=item.row_map_name)
        result, code, output = ran[:3]
        assert code == 0, (result["failure"], result.get("stderr_tail"))
        store.candidates[item.output_blob] = output
        store.candidates[item.row_map_blob] = ran[3]
        store.candidates[item.result_blob] = _json_bytes({**result, "contract_version": "remote-worker.v3",
            "input": {**result["input"], "verified": True}, "parser": {**result["parser"], "verified": True},
            "output": {**result["output"], "content_type": OUTPUT_CONTENT_TYPE},
            "row_map": {**result["row_map"], "content_type": "application/json"}})
        return StartResult(item.job_id, "offline-child-process", datetime.now(timezone.utc))


def published_rows(store, result):
    raw = store.download_output(result["output"]["blob_path"], 1_000_000)
    book = load_workbook(io.BytesIO(raw), read_only=True)
    values = list(book.active.values)
    book.close()
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_a_legacy_xls_combined_bom_with_drawing_rows_resolves_to_one_result_per_board(monkeypatch):
    store = InMemoryStorage()
    source = legacy_xls({"BOM": BOM, "MFG": MFG})
    path = f"customer/rfq-{RFQ}/package-1/original-boms/attachment-1.xls"
    store.sources[path] = source
    documents = [{"document_id": "doc-1", "blob_path": path, "sha256": hashlib.sha256(source).hexdigest(),
                  "size_bytes": len(source)}]
    store.registry[f"candidates/requests/{REQUEST}.json"] = _json_bytes({
        "version": "candidate-request.v2", "request_id": REQUEST, "status": "requested", "parser_key": "Parser1",
        "customer_key": "customer", "profile_hash": "a" * 64, "profile": {"extension": ".xls"},
        "source_evidence": documents})

    # The author is told up front to read a legacy .xls with xlrd, and how parents work in a combined BOM.
    model = Model()
    authored = CandidateAuthoringService(REQUEST, store, model).author()
    assert any(rule.startswith("the input is a legacy binary Excel workbook") for rule in model.rules)
    # A047: values come from sheet.row_values / cell_value / cell.value, never xlrd Cell objects.
    assert any("never an xlrd Cell object" in rule and "sheet.row_values(row)" in rule for rule in model.rules)
    assert any("each assembly's own row has \"line\": null" in rule for rule in model.rules)

    # No association: the Document rows name four items, so there is no single identity and the
    # output still validates; identity is left to review.
    validation = CandidateValidationService(authored.request_id, authored.candidate_id, store,
                                            LocalWorker(store)).validate()
    assert validation.status == "awaiting_approval", validation

    packages = ParentPackageService(store)
    plan = packages.plan(ProcessingRequest.from_dict({"job_id": "combined-legacy", "customer_key": "customer",
        "parser_key": "Parser1", "rfq_uuid": RFQ, "package_version": "1", "documents": documents}),
        approved_assignments=[], pending_assignments=[{"request_id": validation.request_id, "document_ids": ["doc-1"]}])
    choice = {"version": bom_package_choice.RESULT_VERSION_V2, "request_id": "combined-choice", "model_calls": 0,
              "requested_assemblies": REQUESTED, "findings": [], "choices": [], "primary_selections": [],
              "survey": {"files": [], "findings": [], "has_blocking_findings": False},
              "documents": [{"document_id": "doc-1", "sha256": documents[0]["sha256"], "decision": "process",
                             "reason": "combined_bom", "disposition": None, "survey_file_ids": [], "roles": ["bom"],
                             "assemblies": [], "requested_assembly_ids": [A1, A2]}]}
    original = packages._plan
    state = {"request_id": "combined-choice", "status": "completed", "result_sha256": "c" * 64, "result": choice}
    monkeypatch.setattr(packages, "_plan", lambda plan_id: {**original(plan_id),
                                                            "package_choice": bom_package_choice.choice_reference(state)})
    monkeypatch.setattr(bom_package_choice, "bound_state", lambda storage, reference: state)

    review = packages.review(plan["plan_id"])
    assert review["approval_ready"], review
    groups = [group for group in review["groups"] if group.get("status") == "ready"]
    assert sorted(group["evidence"]["requested_link"]["assembly_id"] for group in groups) == [A1, A2]
    approval = packages.approve(plan_id=plan["plan_id"], review_manifest_sha256=review["manifest_sha256"],
                                approver="offline-test")
    resolved = packages.resolve(plan_id=plan["plan_id"], approval_id=approval["approval_id"])
    by_board = {result["assembly_identity"]["name"]: published_rows(store, result) for result in resolved["results"]}
    assert set(by_board) == {"0100-04931", "0100-04932"}
    # Each board keeps only its own lines: its drawings as Documents and its own capacitor with the
    # manufacturer part number from the other sheet.
    assert sorted((row["item_number"], row["item_category"]) for row in by_board["0100-04931"]) == [
        ("0100-04931", "Part"), ("0110-04455", "Document"), ("0120-05191", "Document"), ("0630-03236", "Part")]
    assert sorted((row["item_number"], row["item_category"]) for row in by_board["0100-04932"]) == [
        ("0100-04932", "Part"), ("0110-04456", "Document"), ("0120-05192", "Document"), ("0630-03323", "Part")]
    caps = {row["item_number"]: row["manufacturer_item_number_1"].upper() for rows in by_board.values() for row in rows
            if row["item_number"].startswith("0630-")}
    assert caps == {"0630-03236": "C0805C330J5GACTU", "0630-03323": "C2225C102KGRACTU"}

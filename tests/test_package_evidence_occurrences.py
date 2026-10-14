"""Synthetic resolved-workbook verification and exact row identity fixtures."""
import io
import re
from datetime import datetime
from decimal import Decimal
import zipfile

import pytest
from openpyxl import Workbook

from bom_processor.core import CANONICAL_COLUMNS
from bom_processor.package_evidence_contracts import PackageEvidenceError, digest
from bom_processor.package_evidence_occurrences import quantity_text, read_occurrences

U = "12345678-1234-4234-8234-123456789012"
D = "12345678-1234-4234-8234-123456789014"


def workbook(rows=None, *, sheet="BOM Ω", header=None, extra_sheet=False):
    output = Workbook()
    page = output.active
    page.title = sheet
    page.append(CANONICAL_COLUMNS if header is None else header)
    for changes in rows if rows is not None else [{}]:
        if changes is None:
            page.append([None] * len(CANONICAL_COLUMNS))
            continue
        row = {"line_number": sum(1 for cell in page["A"] if cell.value is not None),
               "item_number": "pcb-01", "item_name": "PCB", "revision": " a ",
               "item_category": "Part", "quantity": 2.500, "reference_designator": "U1,U2"}
        row.update(changes)
        page.append([row.get(column) for column in CANONICAL_COLUMNS])
    if extra_sheet:
        second = output.create_sheet("Other")
        second.append(CANONICAL_COLUMNS)
        second.append([1, "other", "Other", None, "Part", None, 1])
    stream = io.BytesIO()
    output.save(stream)
    return stream.getvalue()


def assembly(data, *, assembly_uuid=U, document_id=D):
    return {"assemblyUuid": assembly_uuid, "partNumber": "ASM", "revision": None,
            "canonicalBom": {"documentId": document_id, "sha256": digest(data), "sizeBytes": len(data)}}


def read(data, **kwargs):
    return read_occurrences([assembly(data)], {D: data}, **kwargs)


def rewrite(data, transform):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            target.writestr(info.filename, transform(info.filename, source.read(info)))
    return output.getvalue()


def test_actual_locators_duplicates_and_document_rows_remain_separate():
    data = workbook([{}, None, {}, {"item_category": "Document", "item_number": "ASM", "quantity": None}])
    result = read(data)
    parts = sorted(result["occurrences"], key=lambda item: item["row"])
    assert [item["row"] for item in parts] == [2, 4]
    assert parts[0]["occurrenceId"] != parts[1]["occurrenceId"]
    assert parts[0]["semanticKey"] == parts[1]["semanticKey"]
    assert parts[0]["partNumber"] == "pcb-01" and parts[0]["revision"] == " a "
    assert parts[0]["quantity"] == "2.5"
    assert parts[0]["sheet"] == "BOM Ω"
    document, = result["documentRows"]
    assert document["row"] == 5 and "occurrenceId" not in document
    assert document["quantity"] is None


@pytest.mark.parametrize("value,expected", [(None, None), (0, "0"), (-0.0, "0"),
    (2.5, "2.5"), ("1.200", "1.2"), ("1e-3", "0.001"), (Decimal("20.00"), "20")])
def test_quantity_text(value, expected):
    assert quantity_text(value) == expected


@pytest.mark.parametrize("value", [True, False, -1, float("nan"), float("inf"), "NaN",
    "Infinity", "1e99999", "-0.1", "", " 2", "bad", datetime(2020, 1, 1), "1e-129"])
def test_invalid_quantities(value):
    with pytest.raises(PackageEvidenceError):
        quantity_text(value)


@pytest.mark.parametrize("changes", [{"item_category": "Other"}, {"item_number": 123},
    {"item_number": None}, {"line_number": 2}, {"line_number": True}, {"quantity": -1},
    {"quantity": True}, {"revision": 3}, {"item_name": None}, {"manufacturer_1": 5},
    {"item_number": datetime(2020, 1, 1)}, {"item_category": None}, {"bom_notes": "=1+1"}])
def test_every_nonempty_row_validated(changes):
    with pytest.raises(PackageEvidenceError):
        read(workbook([changes]))


def test_header_empty_formula_external_macro_and_xml_preflight():
    for data in (workbook(header=["wrong"]), workbook([]), workbook([{"quantity": "=1+1"}])):
        with pytest.raises(PackageEvidenceError):
            read(data)
    base = workbook()
    external = rewrite(base, lambda name, payload: payload.replace(
        b"<Relationships ", b"<Relationships ") if not name.endswith(".rels") else
        payload.replace(b"</Relationships>", b'<Relationship Id="evil" Type="external" '
                        b'Target="https://invalid.example" TargetMode="External"/></Relationships>'))
    with pytest.raises(PackageEvidenceError, match="EXTERNAL_LINK"):
        read(external)
    macro = rewrite(base, lambda name, payload: payload.replace(
        b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")
        if name == "[Content_Types].xml" else payload)
    with pytest.raises(PackageEvidenceError, match="ACTIVE_CONTENT"):
        read(macro)
    entity = rewrite(base, lambda name, payload: b'<!DOCTYPE x [<!ENTITY e "evil">]>' + payload
                     if name == "xl/workbook.xml" else payload)
    with pytest.raises(PackageEvidenceError, match="XML_UNSAFE"):
        read(entity)


def test_workbook_hash_coverage_and_caps():
    data = workbook([{}, {}], extra_sheet=True)
    with pytest.raises(PackageEvidenceError, match="SHEET_LIMIT"):
        read(data, max_sheets=1)
    with pytest.raises(PackageEvidenceError, match="ROW_LIMIT"):
        read(data, max_rows=1)
    supplied = assembly(data)
    supplied["canonicalBom"]["sha256"] = "0" * 64
    with pytest.raises(PackageEvidenceError, match="IDENTITY_MISMATCH"):
        read_occurrences([supplied], {D: data})
    with pytest.raises(PackageEvidenceError, match="COVERAGE"):
        read_occurrences([assembly(data)], {})
    with pytest.raises(PackageEvidenceError, match="COVERAGE"):
        read_occurrences([assembly(data)], {D: data, U: data})


def test_same_workbook_two_assemblies_and_missing_bom_is_explicit():
    data = workbook()
    other = "12345678-1234-4234-8234-123456789015"
    rows = read_occurrences([assembly(data), assembly(data, assembly_uuid=other)], {D: data})
    assert len(rows["occurrences"]) == 2
    assert len({item["occurrenceId"] for item in rows["occurrences"]}) == 2
    missing = assembly(data)
    missing["canonicalBom"] = None
    assert read_occurrences([missing], {}) == {"occurrences": [], "documentRows": [],
                                              "rowObservations": [], "verifiedRows": []}


@pytest.mark.parametrize("old,new", [(b'r="A1"', b'r="XFD1"'),
    (b'<row r="2">', b'<row r="100001">'), (b'r="B2"', b'r="A2"')])
def test_xml_locators_bounded_before_openpyxl(old, new):
    data = rewrite(workbook(), lambda name, payload: payload.replace(old, new)
                   if name == "xl/worksheets/sheet1.xml" else payload)
    with pytest.raises(PackageEvidenceError):
        read(data)


def test_malformed_document_row_not_silently_omitted():
    with pytest.raises(PackageEvidenceError):
        read(workbook([{"item_category": "Document", "item_number": 12}]))


@pytest.mark.parametrize("encoding,bom", [("utf-16-le", b""), ("utf-16-be", b""),
    ("utf-16-le", b"\xff\xfe"), ("utf-16-be", b"\xfe\xff"),
    ("utf-8", b""), ("utf-8", b"\xef\xbb\xbf")])
@pytest.mark.parametrize("part", ["xl/workbook.xml", "xl/worksheets/sheet1.xml",
                                  "docProps/core.xml", "_rels/.rels"])
def test_encoded_internal_entities_rejected_before_parsers(monkeypatch, encoding, bom, part):
    def forbidden(*args, **kwargs):
        pytest.fail("unsafe XML reached a parser")
    data = rewrite(workbook(), lambda name, payload: bom + (
        '<?xml version="1.0" encoding="' + encoding + '"?>'
        '<!DOCTYPE x [<!ENTITY e "opaque">]><x>&e;</x>').encode(encoding)
        if name == part else payload)
    # Earlier safe parts may be parsed, but the unsafe part must never be.
    from bom_processor import package_evidence_occurrences as module
    original = module.ET.fromstring
    def checked(payload):
        if isinstance(payload, bytes):
            decoded = payload[len(bom):].decode(encoding, errors="ignore")
        else:
            decoded = payload
        assert "<!doctype" not in decoded.lower()
        return original(payload)
    monkeypatch.setattr(module.ET, "fromstring", checked)
    monkeypatch.setattr(module.openpyxl, "load_workbook", forbidden)
    with pytest.raises(PackageEvidenceError, match="XML_UNSAFE"):
        read(data)


@pytest.mark.parametrize("encoding,bom", [("utf-16-le", b""), ("utf-16-be", b""),
    ("utf-16-le", b"\xff\xfe"), ("utf-16-be", b"\xfe\xff"),
    ("utf-8", b"\xef\xbb\xbf")])
def test_safe_supported_xml_encodings(encoding, bom):
    def encode(name, payload):
        if name.endswith((".xml", ".rels")):
            content = re.sub(r"^<\?xml\b.*?\?>", "", payload.decode("utf-8"))
            declared = "UTF-16" if encoding.startswith("utf-16") else "UTF-8"
            return bom + ('<?xml version="1.0" encoding="' + declared + '"?>'
                          + content).encode(encoding)
        return payload
    assert len(read(rewrite(workbook(), encode))["occurrences"]) == 1


@pytest.mark.parametrize("payload", [
    b'<?xml version="1.0" encoding="iso-8859-1"?><x/>',
    b'<?xml version="1.0" encoding="utf-16"?><x/>',
    b'\xff\xfe<\x00x\x00/\x00>\x00\xff', b'<x>\xff</x>',
    '<?xml version="1.0"?><x/>'.encode("utf-32"),
    b'<?xml version="1.0" encoding="utf-8?><x/>',
])
def test_unsupported_or_malformed_xml_encodings_fail_before_loading(monkeypatch, payload):
    from bom_processor import package_evidence_occurrences as module
    def forbidden(*args, **kwargs):
        pytest.fail("invalid encoding reached workbook loading")
    monkeypatch.setattr(module.openpyxl, "load_workbook", forbidden)
    data = rewrite(workbook(), lambda name, original: payload
                   if name == "xl/workbook.xml" else original)
    with pytest.raises(PackageEvidenceError, match="XML_(ENCODING|INVALID)"):
        read(data)


def test_complete_row_observations_include_all_manufacturer_slots_and_document_text():
    fields = {"item_name": "Opaque engineering description", "bom_notes": "Opaque notes",
              "unit_of_measure": "EA", "manufacturer_1": "Maker", "manufacturer_item_number_1": "MPN",
              "manufacturer_8": "Alternate", "manufacturer_item_number_8": "MPN8"}
    data = workbook([fields, None, {**fields, "item_category": "Document", "item_name": "Drawing description"}])
    result = read(data)
    observations = result["rowObservations"]
    assert [item["row"] for item in observations] == [2, 4]
    for item, category in zip(observations, ["Part", "Document"]):
        assert set(item) == {"assemblyUuid", "bomDocumentId", "bomSha256", "sheet", "row",
                             "itemCategory", "partNumber", "revision", "itemName", "bomNotes",
                             "unitOfMeasure", "manufacturerPairs"}
        assert item["assemblyUuid"] == U and item["bomDocumentId"] == D
        assert item["bomSha256"] == digest(data) and item["sheet"] == "BOM Ω"
        assert item["itemCategory"] == category and item["partNumber"] == "pcb-01"
        assert item["revision"] == " a " and item["unitOfMeasure"] == "EA"
        assert item["bomNotes"] == fields["bom_notes"]
        assert item["manufacturerPairs"] == [{f"manufacturer_{i}": fields.get(f"manufacturer_{i}"),
            f"manufacturer_item_number_{i}": fields.get(f"manufacturer_item_number_{i}")}
            for i in range(1, 9)]
    assert observations[0]["itemName"] == fields["item_name"]
    assert observations[1]["itemName"] == "Drawing description"


@pytest.mark.parametrize("field,maximum", [("item_number", 256), ("revision", 256),
    ("item_name", 2000), ("bom_notes", 2000), ("unit_of_measure", 2000),
    ("reference_designator", 2000), ("manufacturer_1", 2000),
    ("manufacturer_item_number_8", 2000)])
def test_workbook_text_utf16_exact_boundaries(field, maximum):
    accepted = "😀" * (maximum // 2)
    assert read(workbook([{field: accepted}]))["occurrences"]
    with pytest.raises(PackageEvidenceError, match="INVALID_TEXT"):
        read(workbook([{field: accepted + "x"}]))
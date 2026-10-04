"""Actual offline parser subprocess to package resolution on synthetic XLSX.

Input: synthetic workbooks and injected model response; output: verified resolved
XLSX objects. No Azure/model/CQ calls. Run pytest on this file from repository root.
The test parser supports this fixture layout only, not arbitrary customer BOMs.
"""
import hashlib
import io
import json
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from bom_processor import ParentPackageService, ProcessingRequest
from bom_processor.candidate_authoring import CandidateAuthoringService
from bom_processor.candidate_validation import CandidateValidationService
from bom_processor.core import CANONICAL_COLUMNS, OUTPUT_CONTENT_TYPE, _json_bytes
from bom_processor.foundry_author import FoundryModelAuthor
from bom_processor.job_orchestrator import StartResult
from worker.runtime import Limits, run_job
from test_source_assembly_association import setup


PARSER = '''import argparse
from openpyxl import load_workbook, Workbook
p = argparse.ArgumentParser()
p.add_argument("--input", required=True)
p.add_argument("--output", required=True)
args = p.parse_args()
source = load_workbook(args.input, read_only=True, data_only=False)
columns = COLUMNS
records = []
for sheet in source:
    headers = None
    current = None
    for cells in sheet.iter_rows(values_only=True):
        values = list(cells)
        if "Part" in values and "Quantity" in values:
            headers = {str(value): i for i, value in enumerate(values) if value is not None}
            continue
        if headers is None or not any(value is not None for value in values):
            continue
        def get(key):
            index = headers.get(key)
            return values[index] if index is not None and index < len(values) else None
        item, quantity, maker, mpn = get("Part"), get("Quantity"), get("Maker"), get("MPN")
        if item in ("Drawing", "Procedure"):
            record = dict.fromkeys(columns)
            record.update(item_number=str(quantity), item_name=str(item), item_category="Document")
            records.append(record)
        elif item == "Notes":
            continue
        elif item is None and quantity is None and maker is not None and mpn is not None:
            if current is None:
                raise ValueError("alternate without parent")
            slot = next((i for i in range(1,9) if current.get("manufacturer_"+str(i)) is None), None)
            if slot is None:
                raise ValueError("too many alternatives")
            current["manufacturer_"+str(slot)] = str(maker)
            current["manufacturer_item_number_"+str(slot)] = str(mpn)
        elif item is not None and quantity is not None:
            record = dict.fromkeys(columns)
            record.update(item_number=str(item), item_name=str(item), item_category="Part", quantity=quantity,
                          manufacturer_1=maker, manufacturer_item_number_1=mpn)
            records.append(record)
            current = record
        else:
            raise ValueError("unresolved source row")
source.close()
output = Workbook()
sheet = output.active
sheet.append(columns)
for i, record in enumerate(records, 1):
    record["line_number"] = i
    sheet.append([record[column] for column in columns])
output.save(args.output)
output.close()
check = load_workbook(args.output, read_only=True)
if tuple(next(check.active.values)) != tuple(columns) or check.active.max_row != len(records)+1:
    raise ValueError("output shape")
check.close()
'''.replace("columns = COLUMNS", "columns = " + repr(list(CANONICAL_COLUMNS)))


def execute_package(conflict=False):
    store, sources, associations, _ = setup(conflict=conflict)
    request_path = "candidates/requests/xlsx-demo.json"
    request = json.loads(store.registry[request_path])
    request["version"] = "candidate-request.v2"
    store.registry[request_path] = _json_bytes(request)

    class Agent:
        def __init__(self, **kwargs):
            assert kwargs["tools"] == []
        async def run(self, text):
            return SimpleNamespace(output_text=json.dumps({
                "parser_code": PARSER, "assembly_associations": associations}))

    authored = CandidateAuthoringService("xlsx-demo", store,
        FoundryModelAuthor(agent_factory=Agent), source_associations=True).author()
    executions = []

    class LocalWorker:
        def start(self, item):
            result, code, output = run_job(job=item.job_id,
                parser_bytes=store.candidates[item.parser_blob], input_bytes=store.sources[item.source_blob],
                input_name=item.input_name, output_name=item.output_name, output_policy="opaque",
                return_output=True, limits=Limits(timeout_seconds=15))
            assert code == 0, (result["failure"], result.get("stderr_tail"))
            assert result["status"] == "completed" and output is not None
            executions.append(result)
            # Simulate storage transport only; output comes from the real child process.
            store.candidates[item.output_blob] = output
            remote = {**result, "contract_version": "remote-worker.v1",
                      "input": {**result["input"], "verified": True},
                      "parser": {**result["parser"], "verified": True},
                      "output": {**result["output"], "content_type": OUTPUT_CONTENT_TYPE}}
            store.candidates[item.result_blob] = _json_bytes(remote)
            return StartResult(item.job_id, "offline-child-process", datetime.now(timezone.utc))

    validation = CandidateValidationService(authored.request_id, authored.candidate_id, store, LocalWorker()).validate()
    assert validation.status == "awaiting_approval"
    processing = ProcessingRequest.from_dict({"job_id": "real-worker-demo", "customer_key": "customer",
        "parser_key": "Parser1", "rfq_uuid": "12345678-1234-4234-8234-123456789012",
        "package_version": "1", "documents": sources})
    packages = ParentPackageService(store)
    plan = packages.plan(processing, approved_assignments=[], pending_assignments=[{
        "request_id": authored.request_id, "document_ids": [item["document_id"] for item in sources]}])
    from bom_processor import PackageResolutionError
    try:
        review = packages.review(plan["plan_id"])
    except PackageResolutionError as exc:
        if not conflict or str(exc) != "PACKAGE_CONFLICT_DISPOSITION_REQUIRED":
            raise
        review = {"exception": str(exc)}
    return store, packages, plan, review, executions


def test_real_parser_preserves_alternatives_documents_and_resolves_two_assemblies():
    store, packages, plan, review, executions = execute_package()
    assert len(executions) == 3  # Equivalent source aliases share one execution.
    assert review["approval_ready"] and len(review["groups"]) == 2
    approval = packages.approve(plan_id=plan["plan_id"], review_manifest_sha256=review["manifest_sha256"],
                                 approver="offline-test")
    resolved = packages.resolve(plan_id=plan["plan_id"], approval_id=approval["approval_id"])
    assert len(resolved["results"]) == 2
    for result in resolved["results"]:
        raw = store.download_output(result["output"]["blob_path"], 1_000_000)
        assert hashlib.sha256(raw).hexdigest() == result["output"]["sha256"]
        book = load_workbook(io.BytesIO(raw), read_only=True)
        values = list(book.active.values)
        book.close()
        rows = [dict(zip(values[0], row)) for row in values[1:]]
        docs = [row for row in rows if row["item_category"] == "Document"]
        parts = [row for row in rows if row["item_category"] == "Part"]
        assert {row["item_number"] for row in docs} == {"DRAW-1", "PROC-2"}
        assert all(row["quantity"] is None for row in docs)
        assert len(parts) == 1
        if result["assembly_identity"]["name"] == "asm-100":
            assert parts[0]["quantity"] == 2
            assert parts[0]["manufacturer_item_number_2"] == "m2"
            assert len(result["document_ids"]) == 3
        else:
            assert parts[0]["quantity"] == 4
    assert packages.resolve(plan_id=plan["plan_id"], approval_id=approval["approval_id"]) == resolved


def test_real_parser_conflict_blocks_package_approval_and_publication():
    store, packages, plan, review, executions = execute_package(conflict=True)
    assert len(executions) == 3
    assert review == {"exception": "PACKAGE_CONFLICT_DISPOSITION_REQUIRED"}
    from bom_processor import PackageResolutionError
    with pytest.raises(PackageResolutionError, match="PACKAGE_CONFLICT_DISPOSITION_REQUIRED"):
        packages.approve(plan_id=plan["plan_id"], review_manifest_sha256="a"*64,
                         approver="offline-test")
    assert not any("/resolved/" in path for path in store.outputs)


def test_application_factory_passes_association_policy_to_controller_and_job(monkeypatch):
    import sys
    from types import ModuleType
    from deterministic_api import _candidate_automation_factory
    identity = ModuleType("azure.identity")
    identity.ManagedIdentityCredential = lambda **kwargs: object()
    identity.DefaultAzureCredential = lambda: object()
    monkeypatch.setitem(sys.modules, "azure.identity", identity)
    for key in ("BOM_JOB_SUBSCRIPTION_ID", "BOM_JOB_RESOURCE_GROUP", "BOM_CANDIDATE_AUTHOR_JOB_NAME",
                "BOM_CANDIDATE_VALIDATOR_JOB_NAME", "BOM_CANDIDATE_IDENTITY_CLIENT_ID",
                "BOM_STORAGE_ACCOUNT_URL", "BOM_CANDIDATE_FOUNDRY_PROJECT_ENDPOINT", "BOM_JOB_NAME", "BOM_WORKER_IMAGE"):
        monkeypatch.setenv(key, "offline-test")
    monkeypatch.setenv("BOM_SOURCE_ASSOCIATIONS", "true")
    from bom_processor import InMemoryStorage
    service = _candidate_automation_factory(SimpleNamespace(candidate_automation_enabled=True), InMemoryStorage())
    assert service.authoring_version == "candidate-authoring.v6"
    assert service.author_starter.base_environment["BOM_SOURCE_ASSOCIATIONS"] == "true"
    monkeypatch.setenv("BOM_SOURCE_ASSOCIATIONS", "invalid")
    with pytest.raises(RuntimeError, match="invalid association policy"):
        _candidate_automation_factory(SimpleNamespace(candidate_automation_enabled=True), InMemoryStorage())
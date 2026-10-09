"""Test helper (plan 8.8 step 4a): honest row maps for synthetic worker doubles.

derive_row_map links each populated output value to the first matching source cell, so a
fixture whose output is not really drawn from its source still fails row accounting.
attach_row_map stores it the way the remote worker does (contract remote-worker.v3).
"""
import hashlib
import io
import json
from pathlib import PurePosixPath

from openpyxl import load_workbook

from bom_processor.canonical_format import canonical_pair_count
from bom_processor.core import _json_bytes
from bom_processor.row_accounting import DERIVABLE_FIELDS, VERSION, _same, _text, source_inventory


def _output_rows(output):
    book = load_workbook(io.BytesIO(output), read_only=True, data_only=True)
    rows = [list(row) for row in book.active.iter_rows(values_only=True)]
    book.close()
    header = [str(value) for value in rows[0]]
    assert canonical_pair_count(header) is not None
    return header, rows[1:]


def derive_row_map(source, extension, output):
    original = source_inventory(source, extension)
    header, rows = _output_rows(output)
    by_row = {}
    for (sheet, row, column), value in original.cells.items():
        if _text(value):
            by_row.setdefault((sheet, row), {})[column] = value
    used, lines = set(), []
    for index, values in enumerate(rows, 2):
        fields, cited = {}, set()
        sheet = None
        for field, value in zip(header, values):
            if field in ("line_number", "item_category") or value is None or _text(value) == "":
                continue
            match = next(((s, r, c) for (s, r), cells in sorted(by_row.items())
                          if (sheet is None or s == sheet) and (s, r) not in used
                          for c, v in sorted(cells.items()) if _same(v, value, field == "quantity")), None)
            if match is None:
                if field in DERIVABLE_FIELDS:
                    fields[field] = {"derived": "Synthetic fixture value."}
                continue
            sheet = match[0]
            fields[field] = [match[1], match[2]]
            cited.add(match[1])
        if sheet is None:
            sheet = original.sheets[0]
            cited = {1}
        used.update((sheet, row) for row in cited)
        lines.append({"output_row": index, "sheet": sheet, "rows": sorted(cited), "fields": fields})
    skipped = [{"sheet": sheet, "rows": [[row, row]], "reason": "header" if row == min(r for s, r in by_row if s == sheet) else "other",
                "rationale": "Synthetic fixture row."}
               for (sheet, row) in sorted(by_row) if (sheet, row) not in used]
    line_columns = {(sheet, column) for (sheet, row), cells in by_row.items() if (sheet, row) in used for column in cells}
    cited_columns = {(line["sheet"], link[1]) for line in lines for link in line["fields"].values() if isinstance(link, list)}
    ignored = [{"sheet": sheet, "column": column, "reason": "Synthetic fixture column."}
               for sheet, column in sorted(line_columns - cited_columns)]
    return json.dumps({"version": VERSION, "lines": lines, "skipped": skipped, "ignored_columns": ignored}).encode()


def attach_row_map(storage, item, output, result, source=None):
    """Write the row map and bind it in the worker result when the work item asks for one."""
    if not item.row_map_name:
        return result
    source = source if source is not None else storage.sources[item.source_blob]
    try:
        mapped = derive_row_map(source, PurePosixPath(item.source_blob).suffix.lower(), output)
    except AssertionError:
        # Non-canonical output: the parser still writes some row map; schema validation fails first.
        mapped = json.dumps({"version": VERSION, "lines": [], "skipped": [], "ignored_columns": []}).encode()
    storage.candidates[item.row_map_blob] = mapped
    return {**result, "contract_version": "remote-worker.v3",
            "row_map": {"sha256": hashlib.sha256(mapped).hexdigest(), "size": len(mapped),
                        "content_type": "application/json"}}


def worker_result_bytes(storage, item, output, result):
    return _json_bytes(attach_row_map(storage, item, output, result))

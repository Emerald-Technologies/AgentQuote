"""Sub-assembly structure in the row map (bom-row-map.v2; sub-assembly plan unit 2, Nick 2026-10-12).

The parser states each line's parent and the cell that shows it; code works out the parent from the
cited cells for the four forms seen in customer BOMs. Synthetic data only."""
import copy
import io
import json

import pytest
from openpyxl import Workbook

from bom_processor.canonical_format import canonical_columns
from bom_processor.row_accounting import VERSION, VERSION_V1, RowMapError, verify_row_map

# Part, desc, qty, maker, MPN in columns 2-6; column 1 shows the structure for each form.
PARTS = [("A-1", "Resistor", 1, "MAKER", "A1"), ("SUB-1", "Harness", 2, None, None),
         ("C-1", "Connector", 1, "MAKER", "C1"), ("D-1", "Contact", 10, "MAKER", "D1"),
         ("E-1", "Capacitor", 4, "MAKER", "E1")]
PARENTS = [None, None, 3, 3, None]  # output rows: line n is output row n + 1; SUB-1 is output row 3
FORMS = {
    "level": (["Level"], [1, 1, 2, 2, 1]),
    "number": (["Find"], ["1", "2", "2.1", "2.2", "3"]),
    "parent": (["Parent"], ["TOP-9", "TOP-9", "SUB-1", "SUB-1", "TOP-9"]),
    "indent": (None, None),  # the description is indented instead
}


def xlsx(rows, title="BOM"):
    book = Workbook(); sheet = book.active; sheet.title = title
    for row in rows:
        sheet.append(row)
    data = io.BytesIO(); book.save(data); book.close()
    return data.getvalue()


def build(form, parents=PARENTS, version=VERSION):
    header, marks = FORMS[form]
    rows = [(header or ["Note"]) + ["Part", "Desc", "Qty", "Maker", "MPN"]]
    for index, (part, desc, qty, maker, mpn) in enumerate(PARTS):
        if form == "indent":  # the source always shows the true structure; parents is the claim
            desc = ("    " if PARENTS[index] else "") + desc
        rows.append([None if form == "indent" else marks[index], part, desc, qty, maker, mpn])
    columns = canonical_columns(8)
    out_rows = [columns]
    lines = []
    for index, (part, desc, qty, maker, mpn) in enumerate(PARTS):
        source_row, out = index + 2, index + 2
        values = dict.fromkeys(columns)
        values.update(line_number=index + 1, item_category="Part", item_number=part, item_name=desc, quantity=qty,
                      manufacturer_1=maker, manufacturer_item_number_1=mpn)
        out_rows.append([values[c] for c in columns])
        fields = {"item_number": [source_row, 2], "item_name": [source_row, 3], "quantity": [source_row, 4]}
        if maker:
            fields.update(manufacturer_1=[source_row, 5], manufacturer_item_number_1=[source_row, 6])
        line = {"output_row": out, "sheet": "BOM", "rows": [source_row], "fields": fields}
        if parents is not None:
            line["parent"] = {"line": parents[index], "by": form,
                              "from": [source_row, 3 if form == "indent" else 1]}
        lines.append(line)
    row_map = {"version": version, "lines": lines,
               "skipped": [{"sheet": "BOM", "rows": [[1, 1]], "reason": "header", "rationale": "Headings."}],
               "ignored_columns": []}
    return xlsx(rows), xlsx(out_rows, title="Parsed"), row_map


def run(source, out, row_map):
    return verify_row_map(source, ".xlsx", out, json.dumps(row_map).encode())


@pytest.mark.parametrize("form", sorted(FORMS))
def test_each_form_of_sub_assembly_is_proven_from_the_cited_cells(form):
    result = run(*build(form))
    assert result["valid"] is True, (result["failed_rows"], result["unmapped_columns"])
    assert result["structure"] == {"by": form, "lines_with_parent": 2, "depth": 1}


@pytest.mark.parametrize("form", sorted(FORMS))
def test_a_wrong_parent_claim_fails_on_its_line(form):
    wrong = list(PARENTS); wrong[4] = 3  # E-1 claimed inside the harness
    result = run(*build(form, wrong))
    assert {"sheet": "BOM", "row": 6, "code": "STRUCTURE_PARENT_MISMATCH"} in result["failed_rows"]


def test_quantities_stay_per_parent_and_a_rolled_quantity_fails():
    source, out, row_map = build("level")
    # D-1 sits in a harness used twice: 20 per assembly is code's job, the parser writes 10.
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(out)).active
    sheet.cell(5, canonical_columns(8).index("quantity") + 1).value = 20
    data = io.BytesIO(); sheet.parent.save(data)
    assert "FIELD_VALUE_MISMATCH" in run(source, data.getvalue(), row_map)["failure_codes"]


def test_all_lines_or_none_and_one_kind_and_column():
    source, out, row_map = build("level")
    partial = copy.deepcopy(row_map); del partial["lines"][2]["parent"]
    assert {"sheet": "BOM", "row": 4, "code": "STRUCTURE_INCOMPLETE"} in run(source, out, partial)["failed_rows"]
    mixed = copy.deepcopy(row_map); mixed["lines"][0]["parent"]["by"] = "number"
    assert "STRUCTURE_INCONSISTENT" in run(source, out, mixed)["failure_codes"]


def test_a_cell_that_cannot_show_structure_or_lies_outside_the_line_fails():
    source, out, row_map = build("level")
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(source)); book.active.cell(3, 1).value = "sub"  # not a level
    data = io.BytesIO(); book.save(data)
    result = run(data.getvalue(), out, row_map)
    assert {"sheet": "BOM", "row": 3, "code": "STRUCTURE_VALUE_INVALID"} in result["failed_rows"]
    outside = copy.deepcopy(row_map); outside["lines"][1]["parent"]["from"] = [2, 1]
    assert {"sheet": "BOM", "row": 3, "code": "STRUCTURE_VALUE_INVALID"} in run(source, out, outside)["failed_rows"]


def test_flat_maps_stay_valid_in_both_versions_and_v1_cannot_carry_parents():
    for version in (VERSION_V1, VERSION):
        source, out, row_map = build("level", parents=None, version=version)
        row_map["ignored_columns"] = [{"sheet": "BOM", "column": 1, "reason": "Level."}]
        result = run(source, out, row_map)
        assert result["valid"] is True and result["structure"] == {"by": None, "lines_with_parent": 0, "depth": 0}
    with pytest.raises(RowMapError, match="ROW_MAP_LINE_INVALID"):
        run(*build("level", version=VERSION_V1))


@pytest.mark.parametrize("claim", [{"line": 3, "by": "level"}, {"line": 3, "by": "tree", "from": [2, 1]},
                                   {"line": 1, "by": "level", "from": [2, 1]}, {"line": "3", "by": "level", "from": [2, 1]}])
def test_a_malformed_parent_claim_is_refused(claim):
    source, out, row_map = build("level")
    row_map["lines"][2]["parent"] = claim
    with pytest.raises(RowMapError, match="ROW_MAP_PARENT_INVALID"):
        run(source, out, row_map)


def test_two_levels_deep_and_a_level_jump_use_the_nearest_shallower_line():
    source, out, row_map = build("level")
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(source)); book.active.cell(5, 1).value = 3  # D-1 now under C-1
    data = io.BytesIO(); book.save(data)
    row_map["lines"][3]["parent"]["line"] = 4
    result = run(data.getvalue(), out, row_map)
    assert result["valid"] is True and result["structure"]["depth"] == 2


def test_the_rules_tell_the_parser_how_to_state_parents_and_keep_quantities_as_printed():
    from bom_processor.candidate_authoring import CANONICAL_OUTPUT_RULES, ROW_MAP_RULE
    from bom_processor.document_role import BOM_READING_GUIDANCE
    from bom_processor.row_refinement import refinement_instructions
    assert "bom-row-map.v2" in ROW_MAP_RULE and '"by\\": one of level, number, parent, indent' in json.dumps(ROW_MAP_RULE)
    rules = " ".join(CANONICAL_OUTPUT_RULES)
    assert "never multiply quantities through parents" in rules and "clues read from the data, not rules" in rules
    assert "Sub-assemblies are common" in BOM_READING_GUIDANCE
    assert "STRUCTURE_PARENT_MISMATCH" in refinement_instructions({"rows": []}, 1, 2)

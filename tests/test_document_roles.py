"""Synthetic source-role regressions; run pytest tests/test_document_roles.py."""
import hashlib
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from bom_processor import InMemoryStorage
from bom_processor.candidate_authoring import CandidateAuthoringService
from bom_processor.candidate_validation import CandidateValidationError, CandidateValidationService, _semantic_output
from bom_processor.candidates import persist_candidate
from bom_processor.core import CANONICAL_COLUMNS, _json_bytes
from bom_processor.document_roles import POLICY_VERSION, check_document_roles
from bom_processor.master_registry import profile_bytes
from test_candidate_validation import Model, Orchestrator


def source_bytes(*, quantity=2, ambiguous=False, numeric_document_quantity=None):
    book = Workbook()
    sheet = book.active
    sheet.append(["Assembly", "BOARD-100", "A"])
    sheet.append([])
    sheet.append(["Item", "Description", "Qty", "Manufacturer" if ambiguous else "Manufacturer / Document ID",
                  "MPN" if ambiguous else "MPN / Revision"])
    sheet.append(["R-100", "Resistor", quantity, "Maker One", "RES-10K-0603"])
    sheet.append([None, None, None, "Maker Two", "ALT-10K-0603"])
    sheet.append(["Document reference", "Drawing", numeric_document_quantity, "DRAW-100", "A"])
    sheet.append(["Document reference", "Procedure", None, "PROC-100", "A"])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def output_rows(*, broken=False, quantity=2, numeric_document_quantity=None):
    rows = [dict(line_number=1, item_number="R-100", item_name="Resistor", quantity=quantity,
                 item_category="Part", manufacturer_1="Maker One", manufacturer_item_number_1="RES-10K-0603",
                 manufacturer_2="Maker Two", manufacturer_item_number_2="ALT-10K-0603")]
    for index, identity, name in [(2, "DRAW-100", "Drawing"), (3, "PROC-100", "Procedure")]:
        rows.append(dict(line_number=index, item_number="Document reference" if broken else identity,
                         item_name=name, revision=None if broken else "A", item_category="Document",
                         quantity=numeric_document_quantity if index == 2 else None,
                         manufacturer_1=identity if broken else None,
                         manufacturer_item_number_1="A" if broken else None))
    return rows


def output_bytes(rows):
    book = Workbook()
    sheet = book.active
    sheet.append(list(CANONICAL_COLUMNS))
    for row in rows:
        sheet.append([row.get(column) for column in CANONICAL_COLUMNS])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("broken,valid", [(True, False), (False, True)])
def test_shared_boundary_rejects_wrong_document_roles(broken, valid):
    result = _semantic_output(output_bytes(output_rows(broken=broken)), source_bytes(), "xlsx",
                              infer_document_assembly=False)
    assert result["valid"] is valid
    assert result["document_roles"]["source_document_rows"] == 2
    assert result["document_roles"]["version"] == POLICY_VERSION
    # Role checking does not pretend to cover all component rows.
    assert result["source_rows"] == 0


@pytest.mark.parametrize("mutation", ["drop", "revision", "identity", "name", "quantity", "extra", "category"])
def test_missing_changed_or_extra_document_fails(mutation):
    rows = output_rows()
    if mutation == "drop":
        rows.pop()
    elif mutation == "extra":
        rows.append(dict(rows[-1]))
    elif mutation == "category":
        rows[-1].update(item_category="Part", quantity=1)
    else:
        field = {"revision": "revision", "identity": "item_number", "name": "item_name", "quantity": "quantity"}[mutation]
        rows[-1][field] = 1 if mutation == "quantity" else "changed"
    assert check_document_roles(source_bytes(), rows)["valid"] is False


def test_ambiguous_component_headers_are_not_guessed():
    result = check_document_roles(source_bytes(ambiguous=True), output_rows())
    assert result["valid"] is False
    assert result["result"] == "source_roles_ambiguous"


def test_actual_customer_item_component_headers_use_explicit_document_context():
    book = load_workbook(io.BytesIO(source_bytes()))
    sheet = book.active
    sheet.cell(3, 1, "Customer item")
    sheet.cell(3, 4, "Manufacturer")
    sheet.cell(3, 5, "Manufacturer part number")
    stream = io.BytesIO()
    book.save(stream)
    assert check_document_roles(stream.getvalue(), output_rows())["valid"] is True
    assert check_document_roles(stream.getvalue(), output_rows(broken=True))["valid"] is False


@pytest.mark.parametrize("labels", [("Document reference", "Part"), ("Part", "Document reference")])
def test_duplicate_customer_item_roles_fail_closed(labels):
    book = Workbook()
    book.active.append(["Customer item", "Customer item", "Description", "Manufacturer", "Manufacturer part number"])
    book.active.append([*labels, "Drawing", "DRAW-100", "A"])
    stream = io.BytesIO()
    book.save(stream)
    assert check_document_roles(stream.getvalue(), output_rows())["valid"] is False


def test_failure_feedback_selects_current_validation_not_hash_order():
    from bom_processor.candidate_status import CandidateStatusService
    service, storage, _ = setup_service()
    root = f"candidates/events/{service.request_id}/{service.candidate_id}/"
    for identity, message in [("zzz-old", "stale"), ("aaa-current", "current")]:
        storage.candidates[root + f"validation-failed-{identity}.json"] = _json_bytes({
            "request_id": service.request_id, "candidate_id": service.candidate_id,
            "validation_id": identity, "failure_message": message, "repairable": True})
    selected = CandidateStatusService(storage)._latest_validation_failure(service.request_id, service.candidate_id,
                                                                         "aaa-current")
    assert selected["failure_message"] == "current"
    assert CandidateStatusService(storage)._latest_validation_failure(service.request_id, service.candidate_id,
                                                                      "missing") is None


def test_only_specific_failed_old_layout_policy_is_superseded():
    from bom_processor.validation_policy import superseded_layout_failure
    old = {"version": "candidate-validation.v4", "request_id": "r", "candidate_id": "c",
            "validation_id": "v", "candidate": {"sha256": "a" * 64}, "fixtures": [{"sha256": "b" * 64}],
            "document_role_policy": "xlsx-document-roles.v1", "status": "generated_unvalidated",
           "schema": {"valid": True}, "semantic": {"valid": False, "document_roles": {
               "version": "xlsx-document-roles.v1", "valid": False, "result": "source_roles_ambiguous"}}}
    assert superseded_layout_failure(old)
    assert not superseded_layout_failure({**old, "status": "awaiting_approval"})
    assert not superseded_layout_failure({**old, "schema": {"valid": False}})


def test_two_cell_reference_is_not_hidden_by_preceding_document_table():
    book = Workbook()
    book.active.append(["Document ID", "Description"])
    book.active.append(["PROC-100", "Procedure"])
    book.active.append(["Drawing", "DRAW-100"])
    stream = io.BytesIO()
    book.save(stream)
    rows = [dict(item_category="Document", item_number="PROC-100", item_name="Procedure"),
            dict(item_category="Document", item_number="DRAW-100", item_name="Drawing")]
    assert check_document_roles(stream.getvalue(), rows)["valid"] is True
    assert check_document_roles(stream.getvalue(), rows[:1])["valid"] is False


@pytest.mark.parametrize("revision", [1, "001"])
def test_numeric_revisions_and_leading_zero_text_survive_dataframe_loading(revision):
    source = load_workbook(io.BytesIO(source_bytes()))
    source.active.cell(6, 5, revision)
    source.active.cell(7, 5, revision)
    stream = io.BytesIO()
    source.save(stream)
    rows = output_rows()
    rows[1]["revision"] = revision
    rows[2]["revision"] = revision
    result = _semantic_output(output_bytes(rows), stream.getvalue(), "xlsx", infer_document_assembly=False)
    assert result["valid"] is True


@pytest.mark.parametrize("quantity", [2, 4])
def test_correct_documents_preserve_part_quantity_and_alternates_without_mutation(quantity):
    rows = output_rows(quantity=quantity, numeric_document_quantity=0)
    before = _json_bytes(rows)
    output = output_bytes(rows)
    result = _semantic_output(output, source_bytes(quantity=quantity, numeric_document_quantity=0), "xlsx",
                              infer_document_assembly=False)
    assert result["valid"] is True
    assert _json_bytes(rows) == before
    parsed = list(load_workbook(io.BytesIO(output), read_only=True).active.values)
    part = dict(zip(parsed[0], parsed[1]))
    assert part["quantity"] == quantity
    assert part["manufacturer_item_number_1"] == "RES-10K-0603"
    assert part["manufacturer_item_number_2"] == "ALT-10K-0603"


def setup_service(*, broken=False, grouped=False):
    source = source_bytes()
    paths = ["customer/rfq-test/package-1/original-boms/source.xlsx"]
    if grouped:
        paths.append("customer/rfq-test/package-1/original-boms/alias.xlsx")
    storage = InMemoryStorage(sources={path: source for path in paths})
    request = persist_candidate(storage, "Customer1", "customer-key", profile_bytes(source, ".xlsx"), tuple({
        "document_id": f"doc-{index}", "blob_path": path, "size_bytes": len(source),
        "sha256": hashlib.sha256(source).hexdigest()} for index, path in enumerate(paths)))
    candidate = CandidateAuthoringService(request.request_id, storage, Model()).author()
    worker = Orchestrator(storage, output_bytes(output_rows(broken=broken)))
    digest = hashlib.sha256(source).hexdigest()
    from bom_processor.assembly_association import VERSION
    association = {"version": VERSION, "source_sha256": digest,
                   "name": {"value": "BOARD-100", "cell": {"sheet": "Sheet", "row": 1, "column": 2}},
                   "revision": {"value": "A", "cell": {"sheet": "Sheet", "row": 1, "column": 3}},
                   "rationale": "Explicit assembly declaration", "findings": []}
    return CandidateValidationService(request.request_id, candidate.candidate_id, storage, worker,
                                      assembly_associations={digest: association},
                                      sleep=lambda _: None), storage, worker


@pytest.mark.parametrize("grouped", [False, True])
def test_service_records_repairable_role_feedback_and_never_validates_bad_mapping(grouped):
    service, storage, worker = setup_service(broken=True, grouped=grouped)
    if grouped:
        assert service.validate().status == "generated_unvalidated"
    else:
        with pytest.raises(CandidateValidationError, match="SEMANTIC_INVALID"):
            service.validate()
    assert len(worker.calls) == 1
    assert not any("/validated-" in path for path in storage.candidates)
    events = [json.loads(data) for path, data in storage.candidates.items() if "validation-failed-" in path]
    assert len(events) == 1
    assert events[0]["repairable"] is True
    assert "xlsx_document_roles:document_manufacturer_fields" in json.dumps(events[0])


@pytest.mark.parametrize("grouped", [False, True])
def test_changed_policy_holds_before_worker_and_preserves_old_proof(grouped):
    service, storage, worker = setup_service(grouped=grouped)
    from bom_processor.candidate_validation import _sha256
    request = service._request()
    manifest, _ = service._candidate(request)
    identity = {"version": service.validation_version, "request_id": service.request_id,
                "candidate_id": service.candidate_id, "code_sha256": manifest["code_sha256"]}
    digest = request["source_evidence"][0]["sha256"]
    identity.update({"fixtures": [digest], "assembly_associations": service.assembly_associations} if grouped else
                    {"fixture_sha256": digest, "assembly_association": service._association(
                        storage.sources[request["source_evidence"][0]["blob_path"]], request["source_evidence"][0])})
    old_id = _sha256(_json_bytes(identity))
    old_path = f"candidates/{service.request_id}/{service.candidate_id}/validation/{old_id}/evidence.json"
    old_proof = b'{"historical":"preserve"}'
    storage.candidates[old_path] = old_proof
    before = dict(storage.candidates)
    with pytest.raises(CandidateValidationError, match="VALIDATION_POLICY_REAUTHOR_REQUIRED"):
        service.validate()
    assert storage.candidates[old_path] == old_proof
    assert storage.candidates == before
    assert worker.calls == []


@pytest.mark.parametrize("grouped", [False, True])
def test_current_policy_proof_replays_without_second_worker(grouped):
    service, storage, worker = setup_service(grouped=grouped)
    result = service.validate()
    assert result.status == "awaiting_approval"
    assert service.validate() == result
    assert len(worker.calls) == 1
    assert json.loads(storage.candidates[result.evidence_path])["document_role_policy"] == POLICY_VERSION


def test_failed_v1_layout_is_preserved_and_current_validation_status_is_unambiguous():
    from bom_processor.candidate_status import CandidateStatusService
    service, storage, worker = setup_service()
    first = service.validate()
    evidence = json.loads(storage.candidates[first.evidence_path])
    old = {**evidence, "document_role_policy": "xlsx-document-roles.v1",
           "validation_id": "old-layout", "status": "generated_unvalidated",
           "semantic": {"valid": False, "document_roles": {"version": "xlsx-document-roles.v1",
               "valid": False, "result": "source_roles_ambiguous"}}}
    old_path = first.evidence_path.replace(first.validation_id, "old-layout")
    storage.candidates[old_path] = _json_bytes(old)
    before = storage.candidates[old_path]
    assert service.validate() == first
    status = CandidateStatusService(storage).get(service.request_id)
    assert status["status"] == "awaiting_approval"
    assert status["validation_id"] == first.validation_id
    assert storage.candidates[old_path] == before
    assert len(worker.calls) == 1


def test_current_dispatch_revalidates_old_layout_once_without_author_call():
    from bom_processor.candidate_automation import CandidateAutomationService
    from bom_processor.candidate_status import CandidateStatusService
    from test_serial_validation import planned, Starter
    service, storage, worker = setup_service()
    from bom_processor.author_budget import _membership
    enrolled, _, request_id, _ = planned()
    binding = _membership(enrolled, request_id)
    storage.registry.update(enrolled.registry)
    member = {**binding[0], "request_id": service.request_id}
    storage.registry[f"author-budget/v1/members/{service.request_id}.json"] = _json_bytes(member)
    old = {"request_id": service.request_id, "candidate_id": service.candidate_id,
            "version": "candidate-validation.v4",
            "candidate": {"sha256": service._candidate(service._request())[0]["code_sha256"]},
            "fixtures": service._request()["source_evidence"],
           "document_role_policy": "xlsx-document-roles.v1", "validation_id": "old-layout",
           "status": "generated_unvalidated", "schema": {"valid": True},
           "semantic": {"valid": False, "document_roles": {"version": "xlsx-document-roles.v1",
               "valid": False, "result": "source_roles_ambiguous"}}}
    path = f"candidates/{service.request_id}/{service.candidate_id}/validation/old-layout/evidence.json"
    storage.candidates[path] = _json_bytes(old)
    validator = Starter()
    author = Starter()
    automation = CandidateAutomationService(storage, author, validator)
    status = CandidateStatusService(storage, automation=automation).get(service.request_id)
    assert status["status"] == "generated_unvalidated"
    assert len(validator.calls) == 1 and author.calls == []
    CandidateStatusService(storage, automation=automation).get(service.request_id)
    assert len(validator.calls) == 1
    service.serial_owner = validator.calls[0]["BOM_VALIDATOR_SERIAL_OWNER"]
    result = service.validate()
    assert result.status == "awaiting_approval"
    assert storage.candidates[path] == _json_bytes(old)